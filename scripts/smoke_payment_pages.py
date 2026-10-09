"""Render real local Payment pages and exports using a read-only SQLite connection."""

import argparse
import os
import sys
from io import BytesIO
from pathlib import Path

WORKSPACE = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--database', default='severli.sqlite3')
    args = parser.parse_args()
    path = (WORKSPACE / args.database).resolve(strict=True)
    if path.parent != WORKSPACE or path.suffix != '.sqlite3':
        raise RuntimeError('Gunakan database SQLite lokal dalam workspace.')
    os.environ.update(
        DJANGO_SETTINGS_MODULE='config.settings',
        DEBUG='true',
        DEV_DATABASE_URL='',
        LOCAL_DATABASE=path.as_uri() + '?mode=ro',
        DATABASE_SCHEMA='',
        SECRET_KEY='payment-readonly-smoke-check',
    )
    sys.path.insert(0, str(WORKSPACE))
    import django

    django.setup()
    from django.contrib.messages.storage.fallback import FallbackStorage
    from django.test import RequestFactory
    from openpyxl import load_workbook

    from tracking.models import Invoice, InvoicePayment, User
    from tracking.payment_views import payment_detail, payment_export, payment_list

    factory = RequestFactory()
    ids = list(Invoice.objects.filter(dibatalkan=False).values_list('pk', flat=True))
    expected_payments = InvoicePayment.objects.filter(invoice_id__in=ids).count()
    checked = 0
    for role in User.Role.values:
        user = User(username='read-only-check', role=role, is_active=True)

        def request(url):
            req = factory.get(url)
            req.user = user
            req.session = {}
            req._messages = FallbackStorage(req)
            return req

        response = payment_list(request('/payment/'))
        assert response.status_code == 200, role
        assert b'>Payment</a>' in response.content
        checked += 1
        for pk in ids:
            assert payment_detail(request(f'/payment/{pk}/'), pk).status_code == 200, role
            checked += 1
        response = payment_export(request('/payment/download/'))
        assert response.status_code == 200, role
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ['Rekap Invoice', 'Riwayat Pembayaran']
        assert workbook.worksheets[0].max_row == len(ids) + 5
        assert workbook.worksheets[1].max_row == expected_payments + 5
        workbook.close()
        checked += 1
    print(
        f'{checked} halaman/ekspor lulus pada SQLite read-only; '
        f'{len(ids)} invoice dan {expected_payments} pembayaran.'
    )


if __name__ == '__main__':
    main()
