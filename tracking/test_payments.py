from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest.mock import patch
from uuid import uuid4

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, OperationalError, close_old_connections, connection
from django.db.migrations.executor import MigrationExecutor
from django.db.models import Sum
from django.db.models.deletion import ProtectedError
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .forms import InvoiceForm, PaymentForm
from .models import Invoice, InvoiceAttachment, InvoicePayment, InvoicePo, Log, Master, Po, User
from .money import parse_money
from .payment_services import payment_summary, record_payment
from .services import batalkan_invoice, simpan_invoice

PROOF = b'%PDF-1.4\npayment evidence'


def proof(content=PROOF, name='bukti.pdf'):
    return SimpleUploadedFile(name, content, content_type='application/octet-stream')


def payment_data(amount='500.000.000', kind='cicil', seen='1.000.000.000', **changes):
    return {
        'amount': amount,
        'kind': kind,
        'remaining_seen': seen,
        'payment_date': timezone.localdate(),
        'request_id': uuid4().hex,
        'proof': proof(),
        **changes,
    }


class PaymentTests(TestCase):
    def setUp(self):
        self.buyer = User.objects.create_user('buyer-payment', role='purchasing')
        self.director = User.objects.create_user('director-payment', role='direktur')
        self.admin = User.objects.create_user('admin-payment', role='admin')
        self.accounting = User.objects.create_user('accounting-payment', role='accounting')
        self.other = User.objects.create_user('other-payment', role='other')
        self.today = timezone.localdate()
        self.metadata = {
            'vendor': 'Vendor Payment',
            'nomor': 'INV-PAY-1',
            'tanggal': self.today,
            'total_rp': '1.000.000.000',
            'surat_jalan': '',
            'request_id': uuid4().hex,
        }
        self.groups = [
            {'material': 'Cotton', 'rows': [{'color': 'Cream', 'yards': [Decimal(100)]}]},
            {'material': 'Linen', 'rows': [{'color': 'Black', 'yards': [Decimal(50)]}]},
        ]
        self.invoice = simpan_invoice(self.metadata, self.groups, self.buyer)
        self.attachment = InvoiceAttachment.objects.create(
            invoice=self.invoice,
            filename='invoice.pdf',
            content_type='application/pdf',
            size=len(PROOF),
            content=PROOF,
        )
        self.create_url = reverse('payment_create', args=[self.invoice.pk])

    def pay(self, amount='500.000.000', kind='cicil', **changes):
        return record_payment(self.invoice, payment_data(amount, kind, **changes), self.director)

    def post_data(self, **changes):
        data = payment_data(**changes)
        data['payment_date'] = data['payment_date'].isoformat()
        return data

    def workbook(self, **filters):
        response = self.client.get(reverse('payment_export'), filters)
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))

    def test_new_invoice_and_retries_produce_one_bill_across_multiple_pos(self):
        repeated = simpan_invoice(self.metadata, self.groups, self.buyer)
        self.assertEqual(repeated.pk, self.invoice.pk)
        for index in (1, 2):
            po = Po.objects.create(nomor=f'PO PAYMENT {index}')
            InvoicePo.objects.create(invoice=self.invoice, urut=index, po=po)
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.status, summary.paid, summary.remaining, summary.progress),
            ('Unpaid', 0, Decimal('1000000000.00'), 0),
        )
        self.client.force_login(self.buyer)
        response = self.client.get(reverse('payment_list'))
        self.assertEqual(response.context['page_obj'].paginator.count, 1)
        self.assertNotContains(response, '>Bayar</a>')
        self.assertEqual(Invoice.objects.count(), 1)

    def test_new_invoice_requires_positive_total_and_grouped_money(self):
        for value in ('0', '-1', '', 'NaN', '1.234.56', 1.5):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                simpan_invoice(
                    dict(self.metadata, total_rp=value, request_id=uuid4().hex),
                    self.groups,
                    self.buyer,
                )
        form = InvoiceForm(
            {'vendor': 'Vendor', 'nomor': 'INV', 'tanggal': self.today, 'total_rp': '500.000.000'}
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['total_rp'], Decimal('500000000.00'))
        self.assertEqual(Invoice.objects.count(), 1)

    def test_installments_then_lunas_keep_every_proof_and_correct_positions(self):
        first = self.pay()
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.status, summary.paid, summary.remaining, summary.progress),
            ('Dicicil', 500000000, 500000000, 50),
        )
        second = self.pay('300.000.000', seen='500.000.000', proof=proof(PROOF + b' second'))
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.paid, summary.remaining, summary.progress), (800000000, 200000000, 80)
        )
        third = self.pay('200.000.000', 'lunas', seen='200.000.000', proof=proof(PROOF + b' third'))
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.status, summary.paid, summary.remaining, summary.progress),
            ('Paid', 1000000000, 0, 100),
        )
        self.assertEqual(InvoicePayment.objects.count(), 3)
        for payment, content in (
            (first, PROOF),
            (second, PROOF + b' second'),
            (third, PROOF + b' third'),
        ):
            self.assertEqual(
                bytes(InvoicePayment.objects.get(pk=payment.pk).proof_content), content
            )

    def test_exact_remaining_cicil_becomes_paid(self):
        self.pay('1.000.000.000')
        self.assertEqual(payment_summary(self.invoice).status, 'Paid')

    def test_small_positive_remaining_never_shows_100_percent(self):
        self.pay('999.999.999,99')
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.status, summary.remaining, summary.progress),
            ('Dicicil', Decimal('0.01'), Decimal('99.9')),
        )
        self.client.force_login(self.director)
        self.assertContains(
            self.client.get(reverse('payment_detail', args=[self.invoice.pk])), '99,9%'
        )

    def test_invalid_amounts_do_not_add_payments(self):
        for value in ('0', '-1', '1.000.000.001', '1,001', 'NaN', 'Infinity', 0.1, None):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                self.pay(value)
        self.assertEqual(InvoicePayment.objects.count(), 0)
        self.assertEqual(payment_summary(self.invoice).paid, 0)

    def test_date_kind_request_and_proof_are_required_on_server(self):
        invalid = [
            {'payment_date': None},
            {'payment_date': '2026-10-09'},
            {'payment_date': timezone.now()},
            {'kind': 'convert'},
            {'request_id': ''},
            {'request_id': 'x' * 65},
            {'proof': None},
            {'proof': proof(b'not a pdf')},
            {'proof': proof(PROOF, 'bukti.exe')},
            {'proof': proof(b'%PDF-' + b'x' * (5 * 1024 * 1024))},
        ]
        for values in invalid:
            with self.subTest(values=values.keys()), self.assertRaises(ValidationError):
                self.pay(**values)
        self.assertFalse(InvoicePayment.objects.exists())

    def test_valid_document_types_use_verified_mime_instead_of_browser_mime(self):
        for name, content, mime in (
            ('bukti.pdf', PROOF, 'application/pdf'),
            ('bukti.jpg', b'\xff\xd8\xfftest', 'image/jpeg'),
            ('bukti.jpeg', b'\xff\xd8\xfftest', 'image/jpeg'),
            ('bukti.png', b'\x89PNG\r\n\x1a\ntest', 'image/png'),
        ):
            with self.subTest(name=name):
                payment = self.pay('100', proof=proof(content, name))
                self.assertEqual(payment.proof_content_type, mime)

    def test_unreadable_proof_does_not_change_position(self):
        upload = proof()
        with patch.object(upload.file, 'read', side_effect=OSError('failed upload')):
            with self.assertRaises(ValidationError):
                self.pay(proof=upload)
        self.assertFalse(InvoicePayment.objects.exists())

    def test_all_logged_in_roles_can_view_history_download_and_invoice_attachment(self):
        payment = self.pay()
        urls = [
            reverse('payment_list'),
            reverse('payment_detail', args=[self.invoice.pk]),
            reverse('payment_export'),
            reverse('payment_proof', args=[payment.pk]),
            reverse('invoice_attachment', args=[self.invoice.pk]),
        ]
        for user in (self.buyer, self.director, self.admin, self.accounting, self.other):
            self.client.force_login(user)
            for url in urls:
                with self.subTest(role=user.role, url=url):
                    self.assertEqual(self.client.get(url).status_code, 200)
        self.client.logout()
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 302)

    def test_other_roles_and_superuser_cannot_create_or_upload_payments(self):
        superuser = User.objects.create_superuser('super-payment', password='test-password')
        # Even a superuser whose in-memory role is altered must never inherit payment rights.
        superuser.role = 'direktur'
        for user in (self.buyer, self.admin, self.accounting, self.other, superuser):
            with self.subTest(role=user.role), self.assertRaises(PermissionDenied):
                record_payment(self.invoice, payment_data(), user)
        for user in (self.buyer, self.admin, self.accounting, self.other):
            self.client.force_login(user)
            for method in ('get', 'post'):
                response = getattr(self.client, method)(self.create_url, self.post_data())
                self.assertEqual(response.status_code, 403)
        self.assertFalse(InvoicePayment.objects.exists())

    def test_inactive_director_cannot_record_and_csrf_is_required(self):
        self.director.is_active = False
        with self.assertRaises(PermissionDenied):
            self.pay()
        self.director.is_active = True
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.director)
        self.assertEqual(client.post(self.create_url, self.post_data()).status_code, 403)
        self.assertFalse(InvoicePayment.objects.exists())

    def test_retry_is_idempotent_even_after_paid_and_alternate_money_format(self):
        data = payment_data('1.000.000.000', 'lunas')
        first = record_payment(self.invoice, data, self.director)
        data.update(amount='1000000000.00', proof=proof())
        repeated = record_payment(self.invoice, data, self.director)
        self.assertEqual(first.pk, repeated.pk)
        self.assertEqual(InvoicePayment.objects.count(), 1)
        self.assertEqual(Log.objects.filter(aksi='pembayaran invoice').count(), 1)

    def test_reused_request_with_changed_amount_proof_or_user_is_rejected(self):
        data = payment_data()
        record_payment(self.invoice, data, self.director)
        for changes in (
            {'amount': '100'},
            {'proof': proof(PROOF + b' changed')},
            {'payment_date': self.today - timedelta(days=1)},
        ):
            with self.subTest(changes=changes.keys()), self.assertRaises(ValidationError):
                record_payment(self.invoice, dict(data, **changes), self.director)
        another = User.objects.create_user('director-other', role='direktur')
        with self.assertRaises(ValidationError):
            record_payment(self.invoice, dict(data, proof=proof()), another)
        self.assertEqual(InvoicePayment.objects.count(), 1)

    def test_stale_lunas_rejects_without_changing_agreed_amount(self):
        self.pay()
        with self.assertRaisesMessage(ValidationError, '500.000.000,00'):
            self.pay('1.000.000.000', 'lunas')
        with self.assertRaises(ValidationError):
            self.pay('100', 'lunas', seen='500.000.000')
        self.assertEqual(InvoicePayment.objects.count(), 1)

    def test_stale_lunas_after_another_payment_finishes_invoice_shows_latest_zero(self):
        self.pay('1.000.000.000')
        with self.assertRaisesMessage(ValidationError, 'Sisa tagihan terbaru Rp 0,00'):
            self.pay('1.000.000.000', 'lunas')
        self.assertEqual(InvoicePayment.objects.count(), 1)

    def test_paid_and_cancelled_invoices_reject_direct_extra_payment(self):
        self.pay('1.000.000.000')
        with self.assertRaisesMessage(ValidationError, 'Paid'):
            self.pay('1')
        self.invoice.dibatalkan = True
        self.invoice.save(update_fields=['dibatalkan'])
        with self.assertRaisesMessage(ValidationError, 'dibatalkan'):
            self.pay('1')
        self.assertEqual(InvoicePayment.objects.count(), 1)

    def test_financial_invoice_edits_cancellation_and_delete_are_protected(self):
        self.pay()
        for changes in (
            {'total_rp': '1.000'},
            {'invoice_file': proof()},
            {'nomor': 'CHANGED'},
            {'vendor': 'Changed vendor'},
        ):
            with self.subTest(changes=changes.keys()), self.assertRaises(ValidationError):
                simpan_invoice(
                    dict(self.metadata, request_id=uuid4().hex, **changes),
                    [],
                    self.buyer,
                    self.invoice,
                )
        with self.assertRaises(ValidationError):
            batalkan_invoice(self.invoice, self.buyer)
        with self.assertRaises(ProtectedError):
            self.invoice.delete()
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.total_rp, 1000000000)
        self.assertFalse(self.invoice.dibatalkan)
        self.attachment.refresh_from_db()
        self.assertEqual(bytes(self.attachment.content), PROOF)
        changed = simpan_invoice(
            dict(self.metadata, request_id=uuid4().hex, surat_jalan='SJ nonfinansial'),
            [],
            self.buyer,
            self.invoice,
        )
        self.assertEqual(changed.surat_jalan, 'SJ nonfinansial')

    def test_financial_edit_via_existing_view_is_rejected(self):
        self.pay()
        self.client.force_login(self.buyer)
        response = self.client.post(
            reverse('invoice_edit', args=[self.invoice.pk]),
            {
                'vendor': self.invoice.vendor.name,
                'nomor': self.invoice.nomor,
                'tanggal': self.today.isoformat(),
                'total_rp': '100',
                'existing': '1',
                'group_count': '0',
                'request_id': uuid4().hex,
            },
        )
        self.assertContains(response, 'terkunci')
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.total_rp, 1000000000)

    def test_payments_and_po_completion_remain_independent(self):
        po = Po.objects.create(nomor='PO UNFINISHED', selesai=False)
        InvoicePo.objects.create(invoice=self.invoice, urut=1, po=po)
        self.pay('1.000.000.000')
        po.refresh_from_db()
        self.assertFalse(po.selesai)
        another = Invoice.objects.create(
            vendor=self.invoice.vendor,
            nomor='INV-DONE-PO',
            tanggal=self.today,
            total_rp=100,
            dibuat_oleh=self.buyer,
        )
        done = Po.objects.create(nomor='PO DONE', selesai=True)
        InvoicePo.objects.create(invoice=another, urut=1, po=done)
        self.assertEqual(payment_summary(another).status, 'Unpaid')

    def test_legacy_position_is_unknown_and_cannot_be_inferred_by_edits(self):
        self.invoice.payment_reconciled = False
        self.invoice.save(update_fields=['payment_reconciled'])
        summary = payment_summary(self.invoice)
        self.assertEqual(summary.status, 'Perlu rekonsiliasi')
        self.assertIsNone(summary.paid)
        self.assertIsNone(summary.remaining)
        self.assertIsNone(summary.progress)
        with self.assertRaises(ValidationError):
            self.pay()
        changed = simpan_invoice(
            dict(self.metadata, request_id=uuid4().hex), [], self.buyer, self.invoice
        )
        self.assertFalse(changed.payment_reconciled)
        self.client.force_login(self.director)
        response = self.client.get(reverse('payment_list'), {'status': 'reconciliation'})
        self.assertEqual(response.context['page_obj'].paginator.count, 1)
        self.assertNotContains(response, '>Bayar</a>')
        self.assertEqual(self.client.get(self.create_url).status_code, 302)
        workbook = self.workbook()
        self.assertIsNone(workbook['Rekap Invoice']['F6'].value)
        self.assertEqual(workbook['Rekap Invoice']['I6'].value, 'Perlu rekonsiliasi')
        workbook.close()

    def test_post_conflict_and_invalid_input_are_recoverable(self):
        self.client.force_login(self.director)
        self.pay()
        data = self.post_data(amount='1.000.000.000', kind='lunas')
        response = self.client.post(self.create_url, data)
        self.assertEqual(response.status_code, 409)
        self.assertContains(response, '500.000.000,00', status_code=409)
        self.assertEqual(response.context['form']['amount'].value(), '1.000.000.000')
        data.update(amount='0', proof=proof())
        response = self.client.post(self.create_url, data)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(InvoicePayment.objects.count(), 1)

    def test_form_success_and_retries_do_not_duplicate(self):
        self.client.force_login(self.director)
        data = self.post_data()
        self.assertEqual(self.client.post(self.create_url, data).status_code, 302)
        data['proof'] = proof()
        self.assertEqual(self.client.post(self.create_url, data).status_code, 302)
        self.assertEqual(InvoicePayment.objects.count(), 1)
        payment = InvoicePayment.objects.get()
        response = self.client.get(reverse('payment_proof', args=[payment.pk]))
        self.assertEqual(response.content, PROOF)
        self.assertTrue(response['Content-Disposition'].startswith('inline'))
        response = self.client.get(reverse('payment_proof', args=[payment.pk]), {'download': 1})
        self.assertTrue(response['Content-Disposition'].startswith('attachment'))
        self.assertEqual(response['Cache-Control'], 'private, no-store')

    def test_failed_database_record_rolls_back_proof_and_balance(self):
        self.client.force_login(self.director)
        with patch('tracking.payment_services.log', side_effect=DatabaseError('write failed')):
            response = self.client.post(self.create_url, self.post_data())
        self.assertEqual(response.status_code, 503)
        self.assertFalse(InvoicePayment.objects.exists())
        self.assertEqual(payment_summary(self.invoice).paid, 0)

    def test_read_endpoints_reject_posts_and_sidebar_payment_is_last(self):
        self.client.force_login(self.director)
        for name, args in (
            ('payment_list', []),
            ('payment_detail', [self.invoice.pk]),
            ('payment_export', []),
        ):
            self.assertEqual(self.client.post(reverse(name, args=args)).status_code, 405)
        html = self.client.get(reverse('payment_list')).content.decode()
        nav = html.split('<nav aria-label="Navigasi utama">')[1].split('</nav>')[0]
        self.assertTrue(nav.rstrip().endswith('>Payment</a>'))

    def test_filter_and_export_positions_history_numeric_dates_and_durable_links(self):
        self.pay(payment_date=self.today - timedelta(days=1))
        self.pay('300.000.000', seen='500.000.000')
        Invoice.objects.create(
            vendor=self.invoice.vendor,
            nomor='INV-OTHER',
            tanggal=self.today,
            total_rp=100,
            dibuat_oleh=self.buyer,
        )
        self.client.force_login(self.accounting)
        filters = {
            'q': 'INV-PAY',
            'status': 'dicicil',
            'date_from': self.today.isoformat(),
            'date_to': self.today.isoformat(),
        }
        response = self.client.get(reverse('payment_list'), filters)
        self.assertEqual(response.context['page_obj'].paginator.count, 1)
        self.assertEqual(response.context['page_obj'][0].payment.paid, 800000000)
        workbook = self.workbook(**filters)
        self.assertEqual(workbook.sheetnames, ['Rekap Invoice', 'Riwayat Pembayaran'])
        recap, history = workbook.worksheets
        self.assertEqual((recap.max_row, history.max_row), (6, 7))
        self.assertIn('Periode Invoice', recap['B3'].value)
        self.assertIn('dicicil'.casefold(), recap['B3'].value.casefold())
        self.assertEqual(
            (recap['F6'].value, recap['G6'].value, recap['H6'].value), (800000000, 200000000, 80)
        )
        for cell in ('E6', 'F6', 'G6', 'H6'):
            self.assertEqual(recap[cell].data_type, 'n')
        self.assertEqual(history['H6'].data_type, 'n')
        self.assertEqual(history['E6'].value.date(), self.today - timedelta(days=1))
        self.assertEqual(recap['D6'].value.date(), self.today)
        self.assertIn('/payment/bukti/', history['J6'].hyperlink.target)
        self.assertNotIn('signature=', history['J6'].value)
        self.assertEqual(history['I6'].value, self.director.username)
        workbook.close()

    def test_empty_export_keeps_both_headers_and_invoice_period_is_not_payment_period(self):
        self.pay(payment_date=self.today - timedelta(days=10))
        self.client.force_login(self.buyer)
        for filters in (
            {'q': 'missing'},
            {'status': 'paid'},
            {'date_to': (self.today - timedelta(days=1)).isoformat()},
        ):
            workbook = self.workbook(**filters)
            self.assertEqual([sheet.max_row for sheet in workbook.worksheets], [5, 5])
            self.assertEqual(workbook.worksheets[0]['A5'].value, 'ID invoice')
            self.assertEqual(workbook.worksheets[1]['A5'].value, 'ID transaksi')
            workbook.close()
        response = self.client.get(reverse('payment_list'), {'q': 'Vendor Payment'})
        self.assertEqual(response.context['page_obj'].paginator.count, 1)

    def test_spreadsheet_user_identifiers_cannot_become_formulas(self):
        self.invoice.nomor = '=HYPERLINK("https://example.com")'
        self.invoice.save(update_fields=['nomor'])
        self.pay()
        self.client.force_login(self.accounting)
        workbook = self.workbook()
        self.assertEqual(workbook['Rekap Invoice']['B6'].data_type, 's')
        self.assertEqual(workbook['Riwayat Pembayaran']['C6'].data_type, 's')
        workbook.close()

    def test_invalid_invoice_period_has_visible_error_and_no_broad_export(self):
        self.client.force_login(self.buyer)
        for filters in (
            {'status': 'unknown'},
            {'date_from': 'invalid'},
            {'date_from': '2026-10-10', 'date_to': '2026-10-01'},
        ):
            self.assertEqual(self.client.get(reverse('payment_list'), filters).status_code, 400)
            self.assertEqual(self.client.get(reverse('payment_export'), filters).status_code, 302)

    def test_payment_form_initial_date_money_precision_and_lunas_readonly(self):
        self.client.force_login(self.director)
        response = self.client.get(self.create_url)
        self.assertEqual(response.context['form']['payment_date'].value(), self.today)
        self.assertContains(response, '1.000.000.000,00')
        self.assertContains(response, 'readonly')
        for value in ('500.000.000', '500000000.00', '500.000.000,00'):
            self.assertEqual(parse_money(value), Decimal('500000000.00'))
        form = PaymentForm(self.post_data(), {'proof': proof()})
        self.assertTrue(form.is_valid(), form.errors)

    def test_accounting_login_and_logout_use_payment_destination(self):
        self.accounting.set_password('accounting-password')
        self.accounting.save()
        response = self.client.post(
            reverse('login'),
            {
                'username': self.accounting.username,
                'password': 'accounting-password',
            },
        )
        self.assertRedirects(response, reverse('payment_list'))
        self.assertRedirects(self.client.post(reverse('logout')), reverse('login'))


class PaymentConcurrencyTests(TransactionTestCase):
    def setUp(self):
        self.director = User.objects.create_user('concurrent-director', role='direktur')
        vendor = Master.objects.create(kind='vendor', name='Concurrent Vendor')
        self.invoice = Invoice.objects.create(
            vendor=vendor,
            nomor='CONCURRENT',
            tanggal=timezone.localdate(),
            total_rp=1000,
            dibuat_oleh=self.director,
        )

    def compete(self, payloads):
        barrier = Barrier(2)

        def worker(data):
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                return ('ok', record_payment(self.invoice, data, self.director).pk)
            except ValidationError:
                return ('conflict', None)
            except OperationalError:
                if connection.vendor != 'sqlite':
                    raise
                return ('busy', None)
            finally:
                connection.close()

        with ThreadPoolExecutor(max_workers=2) as pool:
            return list(pool.map(worker, payloads))

    def test_simultaneous_installments_cannot_overpay(self):
        outcomes = self.compete(
            [payment_data('700', seen='1000'), payment_data('700', seen='1000')]
        )
        self.assertEqual(sum(result[0] == 'ok' for result in outcomes), 1, outcomes)
        self.assertEqual(InvoicePayment.objects.aggregate(total=Sum('amount'))['total'], 700)

    def test_simultaneous_same_request_and_retry_create_one_payment(self):
        request_id = uuid4().hex
        outcomes = self.compete(
            [
                payment_data('700', seen='1000', request_id=request_id),
                payment_data('700', seen='1000', request_id=request_id),
            ]
        )
        self.assertTrue(any(result[0] == 'ok' for result in outcomes), outcomes)
        record_payment(
            self.invoice, payment_data('700', seen='1000', request_id=request_id), self.director
        )
        self.assertEqual(InvoicePayment.objects.count(), 1)

    def test_simultaneous_lunas_and_installment_cannot_silently_reduce_lunas(self):
        outcomes = self.compete(
            [payment_data('1000', 'lunas', seen='1000'), payment_data('700', seen='1000')]
        )
        self.assertEqual(sum(result[0] == 'ok' for result in outcomes), 1, outcomes)
        payment = InvoicePayment.objects.get()
        self.assertIn(payment.amount, (Decimal('700'), Decimal('1000')))
        if payment.kind == 'lunas':
            self.assertEqual(payment.amount, 1000)


class PaymentMigrationTests(TransactionTestCase):
    def restore_latest(self):
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())

    def test_migration_preserves_legacy_amounts_attachments_and_marks_unknown_positions(self):
        executor = MigrationExecutor(connection)
        before = [('tracking', '0008_production_sizes')]
        after = [('tracking', '0009_invoice_payments')]
        executor.migrate(before)
        self.addCleanup(self.restore_latest)
        apps = executor.loader.project_state(before).apps
        user = apps.get_model('tracking', 'User').objects.create(username='legacy-payment-user')
        vendor = apps.get_model('tracking', 'Master').objects.create(kind='vendor', name='Legacy')
        old_invoice = apps.get_model('tracking', 'Invoice')
        for number, total in (('LEGACY-ZERO', 0), ('LEGACY-POSITIVE', 1000000000)):
            invoice = old_invoice.objects.create(
                vendor=vendor,
                nomor=number,
                tanggal=timezone.localdate(),
                total_rp=total,
                dibuat_oleh=user,
            )
            apps.get_model('tracking', 'InvoiceAttachment').objects.create(
                invoice=invoice,
                filename='legacy.pdf',
                size=len(PROOF),
                content_type='application/pdf',
                content=PROOF,
            )
        fields = ['pk', 'nomor', 'tanggal', 'total_rp', 'vendor_id', 'dibuat_oleh_id']
        expected = list(old_invoice.objects.order_by('pk').values_list(*fields))
        executor = MigrationExecutor(connection)
        executor.migrate(after)
        current = executor.loader.project_state(after).apps
        invoices = current.get_model('tracking', 'Invoice')
        self.assertEqual(list(invoices.objects.order_by('pk').values_list(*fields)), expected)
        self.assertFalse(invoices.objects.filter(payment_reconciled=True).exists())
        self.assertEqual(current.get_model('tracking', 'InvoicePayment').objects.count(), 0)
        for attachment in current.get_model('tracking', 'InvoiceAttachment').objects.all():
            self.assertEqual(bytes(attachment.content), PROOF)
        new = invoices.objects.create(
            vendor_id=vendor.pk,
            nomor='NEW-AFTER-MIGRATION',
            tanggal=timezone.localdate(),
            total_rp=100,
            dibuat_oleh_id=user.pk,
        )
        self.assertTrue(new.payment_reconciled)
