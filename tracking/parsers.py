import csv
import io
import re
from decimal import Decimal, InvalidOperation
from statistics import median

from openpyxl import load_workbook


def parse_number(token):
    token = token.strip()
    if ',' in token and '.' in token:
        decimal_mark = ',' if token.rfind(',') > token.rfind('.') else '.'
        thousands_mark = '.' if decimal_mark == ',' else ','
        token = token.replace(thousands_mark, '').replace(decimal_mark, '.')
    elif ',' in token:
        token = token.replace(',', '.')
    elif '.' in token and len(token.rsplit('.', 1)[1]) >= 3:
        token = token.replace('.', '')
    if not re.fullmatch(r'\d+(?:\.\d{1,2})?', token):
        raise InvalidOperation
    value = Decimal(token)
    if value <= 0 or value > Decimal('99999999.99'):
        raise InvalidOperation
    return value.quantize(Decimal('0.01'))


def parse_yards(text):
    values, errors = [], []
    lines = text.replace(';', '\n').replace('\t', '\n').splitlines()
    first_content = True
    for line_no, line in enumerate(lines, 1):
        tokens = line.split()
        if not tokens:
            continue
        if (
            first_content
            and any(word in line.lower() for word in ('yard', 'roll', 'jumlah'))
            and all(not char.isdigit() for char in line)
        ):
            first_content = False
            continue
        first_content = False
        for token in tokens:
            try:
                values.append(parse_number(token))
            except InvalidOperation:
                errors.append(f'Baris {line_no}: "{token}" bukan yard yang valid.')
    return values, errors


def suspicious_yards(values):
    if not values:
        return []
    middle = median(values)
    return [
        index + 1
        for index, value in enumerate(values)
        if value > 3 * middle or value < Decimal('0.2') * middle
    ]


def import_yards(upload):
    extension = upload.name.lower().rsplit('.', 1)[-1]
    if extension == 'csv':
        lines = io.StringIO(upload.read().decode('utf-8-sig')).read().splitlines()
        delimiter = ';' if any(';' in line for line in lines[:3]) else ','
        rows = []
        for line in lines:
            try:
                parse_number(line.strip())
                rows.append([line.strip()])
            except InvalidOperation:
                rows.append(next(csv.reader([line], delimiter=delimiter)))
    elif extension == 'xlsx':
        workbook = load_workbook(upload, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
    else:
        return [], ['Pilih file .csv atau .xlsx.']
    values, errors = [], []
    for number, row in enumerate(rows, 1):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        for cell in row:
            if cell is None or str(cell).strip() == '':
                continue
            try:
                values.append(parse_number(str(cell)))
                break
            except InvalidOperation:
                continue
        else:
            if number != 1:
                errors.append(f'Baris {number} tidak memiliki angka yard.')
    if extension == 'xlsx':
        workbook.close()
    return values, errors
