import hashlib
import io
import logging
import uuid
from datetime import timedelta
from decimal import Decimal
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction, connection
from django.db.models import Sum, Q
from django.forms import formset_factory
from django.http import FileResponse, HttpResponse, JsonResponse, Http404
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse as route
from django.utils import timezone
from django.views.decorators.http import require_POST
from . import forms, services, storage
from .models import *
from .presentation import url, cell, row, number, date, STATUS, RESOURCES
from .reports import dataset, TITLES, MODEL_MAP, FILTER_PATHS

def access(request, kind, write=False):
    if kind=='accounts':
        services.require(request.user,['admin'])
    elif kind=='audit':
        services.require(request.user,['admin','purchasing'])
    elif write:
        services.require(request.user,['purchasing'])

def health(request):
    try:
        with connection.cursor() as cur:
            cur.execute('SELECT 1')
        return JsonResponse({'status':'ok'})
    except Exception:
        return JsonResponse({'status':'unavailable'},status=503)

class SignIn(LoginView):
    template_name='login.html'
    authentication_form=forms.LoginForm
    redirect_authenticated_user=True
    def keys(self):
        # Render overwrites client forwarding headers. REMOTE_ADDR is conservative behind other proxies.
        ip=self.request.META.get('REMOTE_ADDR','unknown')
        username=self.request.POST.get('username','').strip().lower()
        return [hashlib.sha256(s.encode()).hexdigest() for s in ['client-account:'+ip+':'+username,'account:'+username]]
    def post(self,request,*args,**kwargs):
        if LoginAttempt.objects.filter(key__in=self.keys(),failures__gte=8,last_at__gte=timezone.now()-timedelta(minutes=15)).exists():
            form=self.get_form()
            form.add_error(None,'Terlalu banyak percobaan masuk. Coba lagi dalam 15 menit.')
            return self.render_to_response(self.get_context_data(form=form),status=429)
        return super().post(request,*args,**kwargs)
    def form_invalid(self,form):
        with transaction.atomic():
            for key in sorted(self.keys()):
                item,_=LoginAttempt.objects.get_or_create(key=key)
                item=LoginAttempt.objects.select_for_update().get(pk=item.pk)
                item.failures=(item.failures if item.last_at>timezone.now()-timedelta(minutes=15) else 0)+1
                item.last_at=timezone.now()
                item.save()
        return super().form_invalid(form)
    def form_valid(self,form):
        LoginAttempt.objects.filter(key=self.keys()[1]).update(failures=0)
        return super().form_valid(form)

@require_POST
def sign_out(request):
    logout(request)
    return redirect('login')

@login_required
def dashboard(request):
    metrics=[]
    physical=services.balance(bucket='physical'); reserved=services.balance(bucket='reserved',location__kind='warehouse')
    buckets=[('Stok fisik',physical,'physical'),('Teralokasi',reserved,'reserved'),('Stok bebas',services.minus(physical,reserved),''),('Dalam perjalanan',services.balance(bucket='transit'),'transit'),('Di CMT',services.balance(bucket='cmt'),'cmt')]
    for label,q,bucket in buckets:
        link='/stock/?scope=warehouse'+('&status='+bucket if bucket else '') if bucket in ['physical','reserved',''] else '/stock/?status='+bucket
        metrics.append({'label':label,'yards':q[1],'rolls':q[0],'link':link})
    orders=Order.objects.exclude(status__in=['closed','cancelled']).select_related('product').order_by('due_date','-id')[:7]
    order_rows=[{'order':o,**services.totals(o)} for o in orders]
    draft_allocations=Allocation.objects.filter(status='draft').count()
    shipments=Shipment.objects.filter(status__in=['dispatched','partially_received']).count()
    exceptions=Discrepancy.objects.filter(resolved=False).count()
    overdue=Order.objects.exclude(status__in=['closed','cancelled','balanced']).filter(due_date__lt=timezone.localdate()).count()
    audit=Audit.objects.select_related('actor').all()
    if request.user.role=='purchasing':
        audit=audit.filter(actor=request.user)
    elif request.user.role=='director':
        audit=Audit.objects.none()
    return render(request,'dashboard.html',{'title':'Ringkasan','active':'dashboard','metrics':metrics,'order_rows':order_rows,'draft_allocations':draft_allocations,'shipment_count':shipments,'exception_count':exceptions,'overdue':overdue,'active_orders':Order.objects.exclude(status__in=['closed','cancelled','draft']).count(),'activities':audit[:5],'incoming':Receipt.objects.filter(status='posted',received_date__year=timezone.localdate().year,received_date__month=timezone.localdate().month).count()})

def list_filters(kind,request):
    form=forms.FilterForm(request.GET)
    for key in list(form.fields):
        if key not in ['q','start','end'] and key not in FILTER_PATHS[kind]:
            form.fields.pop(key)
    if 'status' in form.fields:
        field=FILTER_PATHS[kind]['status']
        choices=['physical','reserved','transit','cmt'] if field=='bucket' else list(STATUS)
        form.fields['status']=__import__('django.forms',fromlist=['ChoiceField']).ChoiceField(label='Status',required=False,choices=[('','Semua status')]+[(x,STATUS[x]) for x in choices])
        if kind=='shipments':
            form.fields['status'].choices += [('awaiting_receipt','Belum selesai diterima')]
    if kind=='stock' and 'order' in form.fields:
        form.fields['order'].label='Lot terkait PO'
        form.fields['product'].label='Lot terkait produk'
    form.is_valid()
    return form

@login_required
def listing(request,kind):
    if kind not in TITLES:
        raise Http404()
    access(request,kind)
    form=list_filters(kind,request)
    options={k:v for k,v in request.GET.items() if not k.startswith('_')}
    if kind not in ['stock','movements']:
        options['_page']=request.GET.get('page',1)
    headers,rows=dataset(kind,form.cleaned_data,request.user,options) if form.is_valid() else ([],[])
    total=options.get('_total',len(rows))
    if '_total' in options:
        page=Paginator(range(total),25).get_page(options['_page_number'])
        page.object_list=rows
    else:
        page=Paginator(rows,25).get_page(request.GET.get('page'))
    params=request.GET.copy(); params.pop('page',None)
    export_allowed=request.user.role in ['purchasing','director','admin']
    create_url=''
    create_label='Tambah'
    if request.user.role=='purchasing' and kind in ['receipts','orders','allocations','masters','progress','finished','warehouse','reconciliation']:
        create_url=route('create',args=[kind]); create_label={'receipts':'Terima bahan','orders':'Buat PO','allocations':'Buat alokasi','masters':'Tambah master','progress':'Catat progres','finished':'Kirim hasil','warehouse':'Terima hasil','reconciliation':'Rekonsiliasi bahan'}[kind]
    elif kind=='accounts' and request.user.role=='admin':
        create_url=route('create',args=[kind]); create_label='Tambah akun'
    tabs=[]
    if kind in ['shipments','cmt','progress','finished','warehouse','exceptions']:
        tabs=[(x,TITLES[x],f'/{x}/') for x in ['shipments','cmt','progress','finished','warehouse','exceptions']]
    if kind in ['stock','movements']:
        tabs=[('stock','Saldo bahan','/stock/'),('movements','Kartu stok','/movements/')]
    if kind=='masters':
        tabs=[(k,label,'/masters/?kind='+k) for k,label in Master.Kind.choices]
    return render(request,'list.html',{'title':TITLES[kind],'active':kind if kind not in ['cmt','progress','finished','warehouse','exceptions'] else 'shipments','kind':kind,'headers':headers,'page':page,'count':total,'filters':form,'params':params.urlencode(),'create_url':create_url,'create_label':create_label,'export_allowed':export_allowed,'tabs':tabs})

@login_required
def export(request,kind):
    if kind not in TITLES:
        raise Http404()
    access(request,kind)
    services.require(request.user,['purchasing','director','admin'])
    form=list_filters(kind,request)
    if not form.is_valid():
        return render(request,'error.html',{'title':'Filter tidak valid','error':form.errors},status=400)
    headers,rows=dataset(kind,form.cleaned_data,request.user,{k:v for k,v in request.GET.items() if not k.startswith('_')})
    if len(rows)>50000:
        return render(request,'error.html',{'title':'Persempit periode laporan','error':'Ekspor maksimal 50.000 baris. Pilih periode atau filter yang lebih spesifik.'},status=400)
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    book=Workbook(); sheet=book.active; sheet.title=TITLES[kind][:31]
    sheet.append(headers or ['Tidak ada data'])
    for item in rows:
        values=[]
        for c in item['cells']:
            value=c['value']
            if c.get('numeric'):
                try:
                    value=Decimal(str(value).replace('.','').replace(',','.'))
                except Exception:
                    value=str(value)
            else:
                value=str(value)+(f" | {c['sub']}" if c.get('sub') else '')
                if value.lstrip().startswith(('=','+','-','@','\t','\r')):
                    value="'"+value
            values.append(value)
        sheet.append(values)
    for c in sheet[1]:
        c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='253D31')
    for cells in sheet.iter_rows(min_row=2):
        for c in cells:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.data_type=='n':
                c.number_format='#,##0.00'
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width=min(55,max(16,max(len(str(c.value or '')) for c in col[:100])+2))
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    meta=book.create_sheet('Filter')
    meta.append(['Laporan',TITLES[kind]]); meta.append(['Waktu ekspor',timezone.localtime().isoformat()]); meta.append(['Jumlah baris',len(rows)])
    for k,v in request.GET.items():
        if k!='page':
            meta.append([k,"'"+str(v)])
    output=io.BytesIO(); book.save(output); output.seek(0)
    services.audit(request.user,'export',request.user,after={'report':kind,'rows':len(rows),'filters':dict(request.GET)})
    return FileResponse(output,as_attachment=True,filename=f'{kind}-{timezone.localdate()}.xlsx',content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@login_required
def reports_home(request):
    return render(request,'reports.html',{'title':'Laporan & arsip','active':'reports','reports':[(k,TITLES[k],f'/{k}/') for k in ['stock','movements','receipts','allocations','cmt','orders','finished','warehouse','reconciliation','exceptions','corrections']]+[('archive','Arsip PO','/orders/?archive=1')]})

def business_data(form):
    return {k:form.cleaned_data[k] for k in form.Meta.fields if k in form.cleaned_data}

def line_data(formset):
    return [{k:form.cleaned_data[k] for k in form.Meta.fields if k in form.cleaned_data} for form in formset if form.cleaned_data and not form.cleaned_data.get('DELETE')]

FORM_MAP={'masters':forms.MasterForm,'receipts':forms.ReceiptForm,'orders':forms.OrderForm,'allocations':forms.AllocationForm,'shipments':forms.ShipmentForm,'cmt':forms.CMTReceiptForm,'progress':forms.ProgressForm,'finished':forms.FinishedForm,'warehouse':forms.WarehouseForm,'reconciliation':forms.ReconciliationForm,'accounts':forms.AccountForm}

@login_required
@transaction.atomic
def edit(request,kind,pk=None):
    if kind not in FORM_MAP:
        raise Http404()
    access(request,kind,True)
    if request.method=='POST':
        try:
            token=uuid.UUID(request.POST.get('token',''))
            expected={'masters':'save_master','receipts':'save_receipt','orders':'save_order','allocations':'save_allocation','shipments':'save_shipment','cmt':'receive_cmt','progress':'save_progress','finished':'send_finished','warehouse':'receive_warehouse','reconciliation':'reconcile','accounts':'save_account'}[kind]
            previous=Operation.objects.filter(key=token,actor=request.user,action=expected,result_id__isnull=False).first()
            if previous:
                return redirect(url(MODEL_MAP[kind].objects.get(pk=previous.result_id)))
        except (ValueError,TypeError):
            pass
    obj=get_object_or_404(MODEL_MAP[kind],pk=pk) if pk else None
    if pk and kind not in ['masters','receipts','orders','allocations','shipments','accounts']:
        raise PermissionDenied('Transaksi posted tidak dapat diedit. Gunakan reversal.')
    if pk and hasattr(obj,'status') and obj.status!='draft':
        raise PermissionDenied('Transaksi sudah terkunci. Gunakan reversal atau koreksi.')
    initial={}
    for field in ['order','cmt','shipment','kind']:
        if request.GET.get(field):
            initial[field]=request.GET[field]
    form_kwargs={'actor':request.user} if kind!='accounts' else {}
    form=FORM_MAP[kind](request.POST or None,request.FILES or None,instance=obj,initial=initial,**form_kwargs)
    formset=None; allocation=None; cmt_line=None
    if kind in ['receipts','allocations','shipments']:
        line_form={'receipts':forms.ReceiptLineForm,'allocations':forms.AllocationLineForm,'shipments':forms.ShipmentLineForm}[kind]
        kwargs={'actor':request.user}
        if kind=='shipments':
            allocation=obj.allocation if obj else get_object_or_404(Allocation,pk=request.GET.get('allocation'))
            kwargs['allocation']=allocation
        factory=formset_factory(line_form,extra=1 if not pk else 0,can_delete=True,max_num=100,validate_max=True,absolute_max=101)
        line_initial=[{f:getattr(line,f) for f in line_form.Meta.fields} for line in obj.lines.all()] if obj else []
        formset=factory(request.POST or None,prefix='lines',initial=line_initial,form_kwargs=kwargs)
    if kind=='cmt':
        cmt_line=get_object_or_404(ShipmentLine,pk=request.GET.get('line'))
    if request.method=='POST':
        valid=form.is_valid()
        lines_valid=formset.is_valid() if formset else True
        if valid and lines_valid:
            try:
                with transaction.atomic():
                    data=business_data(form)
                    token=form.cleaned_data['token']; version=form.cleaned_data.get('version')
                    # Retry must not upload a second copy if the command already completed.
                    evidence=storage.upload(request.user,form.cleaned_data.get('attachment'))
                    if kind in ['receipts','shipments','cmt','progress','finished','warehouse','reconciliation']:
                        data['evidence']=evidence or (getattr(obj,'evidence',None) if obj else None)
                    if kind=='masters':
                        saved=services.save_master(request.user,data,pk=pk,version=version,key=token)
                    elif kind=='receipts':
                        saved=services.save_receipt(request.user,data,line_data(formset),pk=pk,version=version,key=token)
                    elif kind=='orders':
                        saved=services.save_order(request.user,data,pk=pk,version=version,key=token)
                    elif kind=='allocations':
                        saved=services.save_allocation(request.user,data,line_data(formset),pk=pk,version=version,key=token)
                        if request.POST.get('intent')=='allocate':
                            saved=services.post_allocation(request.user,saved.pk,key=uuid.uuid5(token,'post-allocation'))
                    elif kind=='shipments':
                        saved=services.save_shipment(request.user,allocation,data,line_data(formset),pk=pk,version=version,key=token)
                    elif kind=='cmt':
                        saved=services.receive_cmt(request.user,cmt_line,data,key=token)
                    elif kind=='progress':
                        saved=services.save_progress(request.user,data,key=token)
                    elif kind=='finished':
                        saved=services.send_finished(request.user,data,key=token)
                    elif kind=='warehouse':
                        saved=services.receive_warehouse(request.user,data,key=token)
                    elif kind=='reconciliation':
                        saved=services.reconcile(request.user,data,key=token)
                    else:
                        saved=save_account(request.user,data,form.cleaned_data.get('new_password'),pk=pk,key=token)
                    if kind=='allocations' and request.POST.get('intent')=='allocate':
                        messages.success(request,'Bahan berhasil dialokasikan. Stok sudah direservasi.')
                    else:
                        messages.success(request,'Data tersimpan.' if kind in ['masters','accounts','receipts','orders','allocations','shipments'] else 'Laporan berhasil dicatat.')
                    return redirect(url(saved))
            except ValidationError as error:
                form.add_error(None,error)
            except IntegrityError:
                form.add_error(None,'Nomor atau rincian duplikat. Periksa invoice, kode, PO, dan lot yang sama.')
    title=('Ubah ' if pk else 'Tambah ')+TITLES[kind].lower()
    subtitle='Periksa rincian, lalu simpan draft.' if kind in ['receipts','orders','allocations','shipments'] else 'Catat data sesuai laporan yang diterima.'
    if kind=='allocations':
        subtitle='Simpan draft atau langsung alokasikan bahan. Alokasi langsung mereservasi stok yang tersedia.'
    elif kind=='accounts':
        subtitle='Purchasing mengisi seluruh transaksi. Direktur memantau dan mengunduh laporan. Super Admin mengelola akun.'
    if allocation:
        subtitle=f'{allocation} · {allocation.order} · {allocation.cmt.name}'
    if cmt_line:
        q=services.shipment_balance(cmt_line)
        subtitle=f'{cmt_line.shipment.delivery_note} · {cmt_line.allocation_line.lot.code} · Sisa perjalanan {q[0]} roll / {number(q[1],2)} yard'
    response=render(request,'form.html',{'title':title,'active':kind,'kind':kind,'form':form,'formset':formset,'subtitle':subtitle,'back_url':url(obj) if obj else f'/{kind}/','submit_label':'Simpan draft' if kind in ['receipts','orders','allocations','shipments'] else 'Simpan','allocation_editor':kind=='allocations','existing_evidence':getattr(obj,'evidence',None)},status=400 if request.method=='POST' else 200)
    if request.method=='POST':
        # Inline master records and their audits must not survive a failed transaction.
        transaction.set_rollback(True)
    return response

@services.command(['admin'],User)
def save_account(actor,data,password,pk=None):
    users=list(User.objects.select_for_update().filter(Q(role='admin',is_active=True)|Q(pk=pk)).order_by('pk'))
    obj=next((u for u in users if u.pk==pk),None) if pk else None
    if pk and obj is None:
        obj=User.objects.select_for_update().get(pk=pk)
    if not obj:
        obj=User()
    if pk==actor.pk and (data['role']!='admin' or not data['is_active']):
        raise ValidationError('Akun sendiri tidak dapat dinonaktifkan atau diturunkan rolenya.')
    if obj.pk and obj.role=='admin' and (data['role']!='admin' or not data['is_active']) and len([u for u in users if u.role=='admin' and u.is_active])<=1:
        raise ValidationError('Minimal satu Super Admin aktif harus tersedia.')
    before={'role':obj.role,'active':obj.is_active} if obj.pk else {}
    for k,v in data.items():
        setattr(obj,k,v)
    obj.can_adjust=obj.can_reopen=obj.role=='purchasing'
    if password:
        from django.contrib.auth.password_validation import validate_password
        validate_password(password,obj)
        obj.set_password(password)
    elif not obj.pk:
        raise ValidationError('Kata sandi wajib untuk akun baru.')
    obj.is_staff=False; obj.is_superuser=False
    obj.full_clean(); obj.save()
    services.audit(actor,'update_account' if pk else 'create_account',obj,before,{'role':obj.role,'active':obj.is_active,'can_adjust':obj.can_adjust,'can_reopen':obj.can_reopen,'password_changed':bool(password)})
    return obj

def action_url(kind,obj,action):
    return route('action',args=[kind,obj.pk,action])

@login_required
def detail(request,kind,pk):
    if kind not in MODEL_MAP or kind=='audit':
        raise Http404()
    access(request,kind)
    obj=get_object_or_404(MODEL_MAP[kind],pk=pk)
    purchasing=request.user.role=='purchasing'
    actions=[]; sections=[]; metrics=[]
    def add(label,action):
        actions.append({'label':label,'url':action_url(kind,obj,action)})
    def direct(label,link):
        actions.append({'label':label,'url':link})
    if (purchasing and kind in ['masters','receipts','orders','allocations','shipments'] and (not hasattr(obj,'status') or obj.status=='draft')) or (kind=='accounts' and request.user.role=='admin'):
        direct('Ubah',route('edit',args=[kind,pk]))
    if kind=='receipts':
        if purchasing and obj.status=='draft':
            add('Posting penerimaan','post'); add('Batalkan','cancel')
        elif purchasing and obj.status=='posted':
            add('Reversal','reverse')
        sections.append({'title':'Rincian bahan','headers':['Bahan / warna','Lokasi','Roll','Yard','Lot'],'rows':[row(l,[cell(l.material.name,sub=l.color.name),cell(l.warehouse.name),cell(number(l.rolls),numeric=True),cell(number(l.yards,2),numeric=True),cell(l.lot.code,url(l.lot)) if hasattr(l,'lot') else cell('Belum diposting')]) for l in obj.lines.select_related('material','color','warehouse','lot')]})
    elif kind=='allocations':
        if purchasing and obj.status=='draft':
            add('Alokasikan bahan','allocate')
        if purchasing and obj.status in ['allocated','partially_shipped'] and not obj.source_order_id:
            direct('Kirim bahan',f'/shipments/new/?allocation={obj.pk}')
        if purchasing and obj.status not in ['fully_shipped','rejected','cancelled']:
            add('Batalkan alokasi','cancel')
        rows=[]
        for l in obj.lines.select_related('lot','warehouse'):
            free=services.available(l.lot,l.warehouse,obj.source_order)
            rem=services.allocation_remaining(l)
            rows.append(row(l,[cell(l.lot.code,url(l.lot),sub=f'BARIS-{l.pk}'),cell(l.warehouse.name),cell(number(l.rolls),numeric=True),cell(number(l.yards,2),numeric=True),cell(number(free[1],2),sub=f'{free[0]} roll',numeric=True),cell(number(rem[1],2),sub=f'{rem[0]} roll',numeric=True)]))
        sections.append({'title':'Bahan yang dialokasikan','headers':['Lot','Lokasi','Roll','Yard','Bebas saat ini','Sisa reservasi'],'rows':rows})
        decisions=list(obj.decisions.select_related('actor'))
        if decisions:
            sections.append({'title':'Riwayat persetujuan lama','headers':['Keputusan','Pengguna','Tanggal','Catatan'],'rows':[row(d,[cell(d.decision,status=True),cell(d.actor.username),cell(date(timezone.localtime(d.created_at))),cell(d.notes or '—')]) for d in decisions]})
        sections.append({'title':'Pengiriman','headers':['Referensi','Surat jalan','Status'],'rows':[row(s,[cell(str(s),url(s)),cell(s.delivery_note),cell(s.status,status=True)]) for s in obj.shipments.all()]})
    elif kind=='shipments':
        if purchasing and obj.status=='draft':
            add('Posting pengiriman','dispatch')
        if purchasing and obj.status not in ['draft','reversed']:
            add('Reversal pengiriman','reverse')
        rows=[]
        for l in obj.lines.select_related('allocation_line__lot'):
            rem=services.shipment_balance(l)
            rows.append(row(l,[cell(l.allocation_line.lot.code,url(l.allocation_line.lot)),cell(number(l.rolls),numeric=True),cell(number(l.yards,2),numeric=True),cell(number(rem[1],2),sub=f'{rem[0]} roll',numeric=True),cell('Catat penerimaan',f'/cmt/new/?line={l.pk}') if purchasing and obj.status in ['dispatched','partially_received','discrepancy'] else cell('—')]))
        sections.append({'title':'Rincian pengiriman','headers':['Lot','Roll dikirim','Yard dikirim','Sisa perjalanan','Laporan CMT'],'rows':rows})
        sections.append({'title':'Penerimaan CMT','headers':['Tanggal','PIC','Roll','Yard','Kondisi'],'rows':[row(r,[cell(date(r.report_date),url(r)),cell(r.pic),cell(number(r.rolls),numeric=True),cell(number(r.yards,2),numeric=True),cell('reversed' if r.reversed else r.get_condition_display(),status=True)]) for r in CMTReceipt.objects.filter(shipment_line__shipment=obj)]})
    elif kind=='orders':
        t=services.totals(obj)
        metrics=[('Target',obj.target),('Good diterima',t['good']),('Reject',t['reject']),('Sisa target',t['remaining']),('Kelebihan',t['over'])]
        if purchasing:
            if obj.status=='draft':
                add('Aktifkan PO','activate')
            elif obj.status not in ['closed','cancelled']:
                direct('Alokasikan bahan',f'/allocations/new/?order={pk}')
                direct('Catat progres',f'/progress/new/?order={pk}')
                direct('Kirim hasil',f'/finished/new/?order={pk}')
                add('Tutup PO','close')
                if t['over'] and not obj.over_resolved:
                    add('Selesaikan kelebihan hasil','resolve_over')
            if obj.status not in ['closed','cancelled']:
                add('Batalkan PO','cancel')
            elif not obj.archived:
                add('Arsipkan','archive')
            if obj.status=='closed':
                add('Buka kembali','reopen')
        for report_kind in ['allocations','shipments','progress','finished','warehouse','reconciliation']:
            headers,rows=dataset(report_kind,{'order':obj},request.user)
            sections.append({'title':TITLES[report_kind],'headers':headers,'rows':rows})
        transfers=obj.transfer_allocations.select_related('order','cmt').all()
        if transfers.exists():
            sections.append({'title':'Transfer bahan keluar','headers':['Alokasi','PO tujuan','CMT','Status'],'rows':[row(al,[cell(str(al),url(al)),cell(al.order.number,url(al.order)),cell(al.cmt.name),cell(al.status,status=True)]) for al in transfers]})
    elif kind=='stock':
        headers,rows=dataset('movements',{},request.user,{'lot':obj.pk})
        sections.append({'title':'Kartu stok lot','headers':headers,'rows':rows})
        if purchasing:
            direct('Adjustment',f'/adjustment/?lot={obj.pk}')
    elif kind=='finished':
        if purchasing and obj.status!='received' and obj.order.status not in ['closed','cancelled']:
            direct('Terima hasil',f'/warehouse/new/?shipment={pk}')
        sections.append({'title':'Penerimaan gudang','headers':['Tanggal','Good','Reject','Diterima','PIC'],'rows':[row(r,[cell(date(r.report_date),url(r)),cell(number(r.good),numeric=True),cell(number(r.reject),numeric=True),cell(number(r.received),numeric=True),cell(r.pic)]) for r in obj.receipts.filter(reversed=False)]})
    elif kind=='exceptions' and not obj.resolved and purchasing:
        add('Selesaikan selisih','resolve')
    elif kind in ['cmt','warehouse','reconciliation'] and purchasing and not obj.reversed:
        add('Reversal','reverse')
    fields=[]
    if kind=='receipts':
        fields.append(('Referensi invoice',cell(f'INV-{obj.pk}')))
    omit=['id','password','is_staff','is_superuser','can_adjust','can_reopen','version','created_at','updated_at','created_by','evidence','last_login','date_joined','over_resolved']
    for f in obj._meta.fields:
        if f.name in omit:
            continue
        value=getattr(obj,f.name)
        if value in [None,'']:
            continue
        label=forms.LABELS.get(f.name,{'status':'Status','reversed':'Direversal','archived':'Diarsipkan','resolved':'Selesai','resolution':'Penyelesaian','description':'Selisih','reason':'Alasan','allocation':'Alokasi','receipt_line':'Asal bahan','shipment_line':'Baris pengiriman','receipt':'Penerimaan bahan','correction':'Koreksi'}.get(f.name,f.verbose_name))
        if f.is_relation:
            fields.append((label,cell(str(value),url(value))))
        elif f.choices:
            fields.append((label,cell(getattr(obj,'get_'+f.name+'_display')())))
        elif isinstance(value,bool):
            fields.append((label,cell('Ya' if value else 'Tidak')))
        elif isinstance(value,Decimal):
            fields.append((label,cell(number(value,2))))
        elif f.name=='status':
            fields.append((label,cell(value,status=True)))
        else:
            fields.append((label,cell(value)))
    if kind=='orders':
        fields.append(('CMT',cell(', '.join(o.name for o in obj.cmts.all()) or 'Semua CMT aktif')))
    return render(request,'detail.html',{'title':str(obj) if kind in ['receipts','orders','allocations','shipments','finished','stock','masters'] else TITLES[kind],'active':kind,'kind':kind,'object':obj,'fields':fields,'actions':actions,'sections':sections,'metrics':metrics,'evidence':getattr(obj,'evidence',None),'events':Audit.objects.filter(entity=type(obj).__name__,object_id=str(pk)).select_related('actor')[:20] if request.user.role in ['purchasing','admin'] else [],'back_url':f'/{kind}/'})

ACTION_LABELS={'post':'Posting penerimaan','dispatch':'Posting pengiriman','allocate':'Alokasikan bahan','cancel':'Batalkan transaksi','reverse':'Reversal transaksi','activate':'Aktifkan PO','close':'Tutup PO','reopen':'Buka kembali PO','resolve_over':'Selesaikan kelebihan hasil','archive':'Arsipkan PO','resolve':'Selesaikan selisih'}

@login_required
def action(request,kind,pk,action):
    allowed={'receipts':['post','cancel','reverse'],'orders':['activate','cancel','close','reopen','resolve_over','archive'],'allocations':['allocate','cancel'],'shipments':['dispatch','reverse'],'cmt':['reverse'],'warehouse':['reverse'],'reconciliation':['reverse'],'exceptions':['resolve']}
    if action not in allowed.get(kind,[]):
        raise Http404()
    services.require(request.user,['purchasing'])
    obj=get_object_or_404(MODEL_MAP[kind],pk=pk)
    form_class=forms.ResolutionForm if action=='resolve' else forms.ActionForm
    form=form_class(request.POST or None,request.FILES or None)
    if action in ['post','dispatch','allocate','activate','close','archive']:
        form.fields.pop('reason'); form.fields.pop('attachment')
    impact={'post':'Posting menambah stok fisik sesuai seluruh rincian penerimaan. Data posted dikunci.','dispatch':'Posting mengurangi stok fisik dan reservasi; bahan masuk stok dalam perjalanan.','allocate':'Sistem memeriksa stok bebas dan langsung mereservasi bahan untuk PO ini. Rincian alokasi dikunci setelah berhasil. Stok fisik tetap sampai pengiriman dicatat.','cancel':'Transaksi dibatalkan. Sisa reservasi alokasi dilepas; riwayat tetap tersimpan.','reverse':'Dampak transaksi dibalik dengan pergerakan lawan. Reversal ditolak bila stok sudah dipakai.','close':'PO dikunci setelah target good, penerimaan, dan rekonsiliasi bahan selesai.','resolve':'Selisih diselesaikan dengan bukti. Kehilangan mengurangi sisa perjalanan, tanpa menambah good.','resolve_over':'Kelebihan hasil tetap ditampilkan. Alasan ini mengizinkan penutupan setelah syarat lain selesai.','archive':'PO tersimpan di arsip dan tetap dapat ditelusuri.','activate':'PO dapat digunakan untuk alokasi dan produksi.','reopen':'PO dapat menerima transaksi lagi. Alasan dan pengguna akan dicatat.'}[action]
    if request.method=='POST' and form.is_valid():
        try:
            with transaction.atomic():
                data=form.cleaned_data; token=data['token']
                previous=Operation.objects.filter(key=token,actor=request.user,result_id__isnull=False).exists()
                evidence=None if previous else storage.upload(request.user,data.get('attachment')); reason=data.get('reason','')
                if action=='reverse':
                    services.reverse(request.user,{'receipts':'receipt','shipments':'shipment','cmt':'cmt_receipt','warehouse':'warehouse_receipt','reconciliation':'reconciliation'}[kind],pk,reason,evidence,key=token)
                elif kind=='receipts':
                    services.post_receipt(request.user,pk,key=token) if action=='post' else services.cancel_receipt(request.user,pk,reason,key=token)
                elif kind=='orders':
                    services.order_action(request.user,pk,action,reason,evidence,key=token)
                elif kind=='allocations':
                    if action=='allocate':
                        services.post_allocation(request.user,pk,key=token)
                    else:
                        services.cancel_allocation(request.user,pk,reason,key=token)
                elif kind=='shipments':
                    services.dispatch(request.user,pk,key=token)
                else:
                    services.resolve_discrepancy(request.user,pk,data['resolution'],reason,evidence,data['rolls'],data['yards'],key=token)
                messages.success(request,'Tindakan berhasil dicatat.')
                return redirect(url(obj))
        except ValidationError as error:
            form.add_error(None,error)
        except IntegrityError:
            form.add_error(None,'Data berubah atau permintaan duplikat. Muat ulang detail transaksi.')
    summary=[]
    if kind in ['receipts','shipments','allocations']:
        q=obj.lines.aggregate(r=Sum('rolls'),y=Sum('yards'))
        summary=[('Roll',number(q['r'] or 0)),('Yard',number(q['y'] or ZERO,2))]
    return render(request,'form.html',{'title':ACTION_LABELS[action],'subtitle':str(obj),'active':kind,'form':form,'impact':impact,'summary':summary,'back_url':url(obj),'submit_label':ACTION_LABELS[action]},status=400 if request.method=='POST' else 200)

@login_required
def adjustment(request):
    services.require(request.user,['purchasing'])
    form=forms.AdjustmentForm(request.POST or None,request.FILES or None,initial={'lot':request.GET.get('lot')})
    if request.method=='POST' and form.is_valid():
        try:
            with transaction.atomic():
                data=form.cleaned_data
                evidence=storage.upload(request.user,data.get('attachment'))
                obj=services.adjust(request.user,data['lot'],data['warehouse'],data['direction'],data['rolls'],data['yards'],data['reason'],evidence,key=data['token'])
                messages.success(request,'Adjustment dicatat pada kartu stok.')
                return redirect(url(obj))
        except ValidationError as error:
            form.add_error(None,error)
    return render(request,'form.html',{'title':'Adjustment bahan','active':'stock','form':form,'impact':'Adjustment menambah atau mengurangi stok fisik. Alasan, bukti, serta saldo sebelum dan sesudah dicatat.','submit_label':'Posting adjustment','back_url':'/stock/'},status=400 if request.method=='POST' else 200)

@login_required
def download_evidence(request,pk):
    evidence=get_object_or_404(Evidence,pk=pk)
    try:
        return FileResponse(storage.read(evidence),as_attachment=True,filename=evidence.name,content_type=evidence.mime_type)
    except ValidationError as error:
        return render(request,'error.html',{'title':'Lampiran belum tersedia','error':'; '.join(error.messages)},status=503)

def forbidden(request,exception=None):
    return render(request,'error.html',{'title':'Akses dibatasi','error':str(exception) or 'Akun ini tidak memiliki izin untuk tindakan tersebut.'},status=403)

def not_found(request,exception=None):
    return render(request,'error.html',{'title':'Data tidak ditemukan','error':'Periksa kembali tautan atau buka daftar transaksi.'},status=404)

def server_error(request):
    logging.getLogger(__name__).error('Server error request_id=%s',getattr(request,'request_id',''))
    return render(request,'error.html',{'title':'Permintaan belum dapat diproses','error':'Data belum berubah jika transaksi gagal. Coba kembali atau sampaikan kode permintaan kepada administrator.'},status=500)
