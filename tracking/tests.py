import io
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date,timedelta
from decimal import Decimal
from unittest import skipUnless
from django.conf import settings
from django.core.exceptions import PermissionDenied,ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection,connections,transaction,DatabaseError,IntegrityError
from django.test import TestCase,TransactionTestCase,Client,override_settings
from django.utils import timezone
from .models import *
from . import services as s,storage
from .views import save_account

D=Decimal

class Fixture:
    def build(self):
        self.p=User.objects.create(username='purchasing',role='purchasing',can_adjust=True,can_reopen=True)
        self.a=User.objects.create(username='approver',role='approver')
        self.m=User.objects.create(username='management',role='management')
        self.admin=User.objects.create(username='admin',role='admin')
        def master(kind,code):
            return s.save_master(self.p,{'kind':kind,'code':code,'name':code})
        self.vendor=master('vendor','VENDOR'); self.cmt=master('cmt','CMT'); self.wh=master('warehouse','WH')
        self.material=master('material','INESA'); self.color=master('color','BLACK'); self.unit=master('unit','YRD'); self.product=master('product','FUJI MIDI')
        self.today=timezone.localdate()
        self.receipt=s.save_receipt(self.p,{'vendor':self.vendor,'invoice':' inv 001 ','warehouse':self.wh,'received_date':self.today,'invoice_total':D('100000.00')},[self.receipt_line()])
        self.order=self.new_order('PO 0111')
        self.evidence=Evidence.objects.create(created_by=self.p,name='bukti.pdf',mime_type='application/pdf',size=20,sha256='f'*64,storage_key='test-proof.pdf')
    def receipt_line(self,rolls=50,yards='1665.50'):
        return {'material':self.material,'color':self.color,'unit':self.unit,'warehouse':self.wh,'rolls':rolls,'yards':D(yards)}
    def new_order(self,number,target=1072):
        o=s.save_order(self.p,{'number':number,'product':self.product,'target':target,'cmts':[self.cmt],'order_date':self.today,'due_date':self.today+timedelta(days=10)})
        s.order_action(self.p,o.pk,'activate'); o.refresh_from_db(); return o
    def posted(self):
        if not Lot.objects.filter(receipt_line__receipt=self.receipt).exists():
            s.post_receipt(self.p,self.receipt.pk)
        self.lot=self.receipt.lines.first().lot
        return self.lot
    def allocate(self,r=10,y='300',order=None,approve=True):
        self.posted()
        al=s.save_allocation(self.p,{'order':order or self.order,'cmt':self.cmt},[{'lot':self.lot,'warehouse':self.wh,'rolls':r,'yards':D(y)}])
        s.submit_allocation(self.p,al.pk)
        if approve: s.decide_allocation(self.a,al.pk,'approved')
        al.refresh_from_db(); return al
    def shipment(self,allocation=None,r=10,y='300',dispatch=True):
        allocation=allocation or self.allocate(r,y)
        sj=s.save_shipment(self.p,allocation,{'delivery_note':'SJ-'+uuid.uuid4().hex[:8],'sent_date':self.today},[{'allocation_line':allocation.lines.first(),'rolls':r,'yards':D(y)}])
        if dispatch: s.dispatch(self.p,sj.pk)
        sj.refresh_from_db(); return sj
    def report(self,**overrides):
        return {'pic':'PIC','medium':'whatsapp','report_date':self.today,'notes':'Laporan diterima',**overrides}
    def received_cmt(self,sj=None,r=10,y='300',**extra):
        sj=sj or self.shipment(r=r,y=y)
        return s.receive_cmt(self.p,sj.lines.first(),self.report(rolls=r,yards=D(y),condition='good',**extra))
    def finished(self,qty,good=None,reject=0,order=None,**extra):
        order=order or self.order
        fs=s.send_finished(self.p,self.report(order=order,cmt=self.cmt,warehouse=self.wh,delivery_note='HSL-'+uuid.uuid4().hex[:8],quantity=qty))
        wr=s.receive_warehouse(self.p,self.report(shipment=fs,received=qty,good=qty-reject if good is None else good,reject=reject,discrepancy=False,**extra))
        return fs,wr

class AcceptanceTests(Fixture,TestCase):
    def setUp(self): self.build()
    def test_at01_post_exact_decimal_and_once(self):
        token=uuid.uuid4(); r=s.post_receipt(self.p,self.receipt.pk,key=token)
        s.post_receipt(self.p,self.receipt.pk,key=token)
        self.assertEqual(s.balance(bucket='physical'),(50,D('1665.50')))
        self.assertEqual(Movement.objects.count(),1)
        self.assertEqual(r.status,'posted')
    def test_at02_normalized_duplicate_invoice(self):
        with self.assertRaises(ValidationError):
            s.save_receipt(self.p,{'vendor':self.vendor,'invoice':'INV   001','warehouse':self.wh},[self.receipt_line()])
        self.assertEqual(Receipt.objects.count(),1); self.assertEqual(Movement.objects.count(),0)
    def test_at03_allocation_exceeding_stock_rejected(self):
        with self.assertRaises(ValidationError): self.allocate(51,'1665.50')
        self.assertEqual(s.balance(bucket='reserved'),(0,D('0')))
    def test_at05_approval_reserves_without_physical_reduction(self):
        self.allocate(10,'300.25')
        self.assertEqual(s.balance(bucket='physical'),(50,D('1665.50')))
        self.assertEqual(s.balance(bucket='reserved'),(10,D('300.25')))
        self.assertEqual(s.available(self.lot,self.wh),(40,D('1365.25')))
    def test_at06_partial_dispatch(self):
        al=self.allocate(); self.shipment(al,4,'120')
        self.assertEqual(s.balance(bucket='physical'),(46,D('1545.50')))
        self.assertEqual(s.balance(bucket='reserved'),(6,D('180')))
        self.assertEqual(s.balance(bucket='transit'),(4,D('120')))
        al.refresh_from_db(); self.assertEqual(al.status,'partially_shipped')
    def test_at07_partial_cmt_receipt(self):
        sj=self.shipment(); self.received_cmt(sj,4,'100')
        self.assertEqual(s.balance(bucket='transit'),(6,D('200')))
        self.assertEqual(s.balance(bucket='cmt'),(4,D('100')))
        sj.refresh_from_db(); self.assertEqual(sj.status,'partially_received')
    def test_at08_discrepancy_requires_evidence_and_resolution(self):
        sj=self.shipment()
        data=self.report(rolls=8,yards=D('240'),condition='discrepancy')
        with self.assertRaises(ValidationError): s.receive_cmt(self.p,sj.lines.first(),data)
        data['evidence']=self.evidence
        s.receive_cmt(self.p,sj.lines.first(),data)
        sj.refresh_from_db(); self.assertEqual(sj.status,'discrepancy')
        d=Discrepancy.objects.get()
        with self.assertRaises(ValidationError): s.resolve_discrepancy(self.p,d.pk,'confirmed','Sudah sesuai',self.evidence)
        s.resolve_discrepancy(self.p,d.pk,'loss','Hilang dalam pengiriman',self.evidence,2,D('60'))
        sj.refresh_from_db(); self.assertEqual(sj.status,'received')
        self.assertEqual(s.balance(bucket='transit'),(0,D('0')))
    def test_at09_good_only_reduces_remaining(self):
        self.allocate(); self.finished(440,good=430,reject=10)
        self.assertEqual(s.totals(self.order)['remaining'],642)
        self.assertEqual(s.totals(self.order)['reject'],10)
    def test_at10_prd_po0111_end_to_end(self):
        self.received_cmt()
        for amount in [430,247,250,145]: self.finished(amount)
        self.order.refresh_from_db(); self.assertEqual(self.order.status,'balance_candidate')
        self.assertEqual(s.totals(self.order)['good'],1072)
        s.reconcile(self.p,self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='consumed',rolls=10,yards=D('300')))
        self.order.refresh_from_db(); self.assertEqual(self.order.status,'balanced')
        s.order_action(self.p,self.order.pk,'close')
        self.order.refresh_from_db(); self.assertEqual(self.order.status,'closed')
    def test_at11_overdelivery_requires_manual_resolution(self):
        al=self.allocate(); s.cancel_allocation(self.p,al.pk,'Tidak memerlukan bahan tambahan')
        self.finished(1080)
        self.assertEqual(s.totals(self.order)['over'],8)
        with self.assertRaises(ValidationError): s.order_action(self.p,self.order.pk,'close')
        s.order_action(self.p,self.order.pk,'resolve_over','8 pcs tambahan diterima')
        s.order_action(self.p,self.order.pk,'close')
    def test_at12_posted_edit_rejected(self):
        self.posted()
        with self.assertRaises(ValidationError): s.save_receipt(self.p,{'vendor':self.vendor,'invoice':'edit','warehouse':self.wh},[],pk=self.receipt.pk,version=1)
        self.assertEqual(s.balance(bucket='physical'),(50,D('1665.50')))
    def test_at13_management_and_admin_api_denied(self):
        for user in [self.m,self.a,self.admin]:
            self.client.force_login(user)
            before=Audit.objects.count()
            response=self.client.post(f'/receipts/{self.receipt.pk}/action/post/',{'token':str(uuid.uuid4())})
            self.assertEqual(response.status_code,403)
            self.assertEqual(Audit.objects.count(),before)
        self.assertEqual(Movement.objects.count(),0)
    def test_at14_export_matches_active_filter_and_no_formula_injection(self):
        self.allocate(); self.finished(20)
        self.client.force_login(self.m)
        from openpyxl import load_workbook
        response=self.client.get('/warehouse/export/',{'cmt':self.cmt.pk,'start':self.today.isoformat(),'end':self.today.isoformat()})
        self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(wb.active.max_row,2)
        self.assertEqual(wb.active.cell(2,5).value,20)
        response=self.client.get('/warehouse/export/',{'start':(self.today+timedelta(days=1)).isoformat()})
        wb=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(wb.active.max_row,1)
    def test_at15_close_unresolved_rejected(self):
        self.received_cmt(); self.finished(1072)
        with self.assertRaises(ValidationError): s.order_action(self.p,self.order.pk,'close')
    def test_warehouse_complete_loss_can_be_reported_without_inventing_receipt(self):
        self.allocate()
        fs=s.send_finished(self.p,self.report(order=self.order,cmt=self.cmt,warehouse=self.wh,delivery_note='HSL-LOST',quantity=440))
        wr=s.receive_warehouse(self.p,self.report(shipment=fs,received=0,good=0,reject=0,discrepancy=True,evidence=self.evidence,notes='Seluruh kiriman belum diterima dan dinyatakan hilang.'))
        discrepancy=Discrepancy.objects.get(finished_shipment=fs)
        s.resolve_discrepancy(self.p,discrepancy.pk,'loss','Kehilangan 440 pcs dikonfirmasi',self.evidence,440,D('0'))
        fs.refresh_from_db()
        self.assertEqual(fs.status,'received')
        self.assertEqual(s.totals(self.order)['good'],0)
        self.assertEqual(s.totals(self.order)['remaining'],1072)
        self.assertEqual(wr.received,0)
    def test_close_waits_for_outgoing_transfer_requests(self):
        self.received_cmt(); target=self.new_order('PO TRANSFER PENDING')
        al=s.save_allocation(self.p,{'order':target,'source_order':self.order,'cmt':self.cmt},[{'lot':self.lot,'warehouse':self.cmt,'rolls':4,'yards':D('120')}])
        self.finished(1072)
        s.reconcile(self.p,self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='consumed',rolls=10,yards=D('300')))
        with self.assertRaises(ValidationError): s.order_action(self.p,self.order.pk,'close')
        s.cancel_allocation(self.p,al.pk,'Permintaan transfer dibatalkan')
        s.order_action(self.p,self.order.pk,'close')
    def test_at16_duplicate_cmt_request(self):
        sj=self.shipment(); token=uuid.uuid4()
        data=self.report(rolls=4,yards=D('120'),condition='good')
        s.receive_cmt(self.p,sj.lines.first(),data,key=token)
        s.receive_cmt(self.p,sj.lines.first(),data,key=token)
        self.assertEqual(CMTReceipt.objects.count(),1)
        self.assertEqual(s.balance(bucket='cmt'),(4,D('120')))
    def test_atomic_multiline_receipt_rollback(self):
        r=s.save_receipt(self.p,{'vendor':self.vendor,'invoice':'SECOND','warehouse':self.wh},[self.receipt_line(),self.receipt_line()])
        self.color.active=False; self.color.save()
        with self.assertRaises(ValidationError): s.post_receipt(self.p,r.pk)
        self.assertFalse(Lot.objects.exists()); self.assertFalse(Movement.objects.exists())
        r.refresh_from_db(); self.assertEqual(r.status,'draft')
    def test_approval_rechecks_stock_after_submit(self):
        first=self.allocate(40,'1300',approve=False)
        second=self.allocate(40,'1300',order=self.new_order('PO SECOND'),approve=False)
        s.decide_allocation(self.a,first.pk,'approved')
        with self.assertRaises(ValidationError): s.decide_allocation(self.a,second.pk,'approved')
        self.assertEqual(Decision.objects.count(),1)
    def test_purchasing_cannot_approve(self):
        al=self.allocate(approve=False)
        with self.assertRaises(PermissionDenied): s.decide_allocation(self.p,al.pk,'approved')
    def test_approval_cannot_be_decided_twice(self):
        al=self.allocate()
        with self.assertRaises(ValidationError): s.decide_allocation(self.a,al.pk,'rejected','Tidak')
        self.assertEqual(Decision.objects.count(),1)
    def test_revision_updates_require_version_and_resubmit(self):
        al=self.allocate(approve=False)
        s.decide_allocation(self.a,al.pk,'revision','Kurangi bahan'); al.refresh_from_db()
        with self.assertRaises(ValidationError): s.save_allocation(self.p,{'order':self.order,'cmt':self.cmt},[],pk=al.pk,version=1)
        al=s.save_allocation(self.p,{'order':self.order,'cmt':self.cmt},[{'lot':self.lot,'warehouse':self.wh,'rolls':2,'yards':D('60')}],pk=al.pk,version=al.version)
        s.submit_allocation(self.p,al.pk); s.decide_allocation(self.a,al.pk,'approved')
        self.assertEqual(s.balance(bucket='reserved'),(2,D('60')))
    def test_cancellation_releases_only_remaining_reservation(self):
        al=self.allocate(); self.shipment(al,4,'120')
        s.cancel_allocation(self.p,al.pk,'Sisa dibatalkan')
        self.assertEqual(s.balance(bucket='reserved'),(0,D('0')))
        self.assertEqual(s.balance(bucket='physical'),(46,D('1545.50')))
    def test_cmt_overreceipt_rejected(self):
        sj=self.shipment()
        with self.assertRaises(ValidationError): self.received_cmt(sj,11,'301')
        self.assertEqual(s.balance(bucket='transit'),(10,D('300')))
    def test_received_good_reject_exact_sum(self):
        self.allocate()
        with self.assertRaises(ValidationError): self.finished(440,good=440,reject=10)
        self.assertFalse(WarehouseReceipt.objects.exists())
    def test_return_and_waste_preserve_lot_and_require_evidence(self):
        self.received_cmt()
        s.reconcile(self.p,self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='returned',warehouse=self.wh,rolls=2,yards=D('60')))
        self.assertEqual(s.balance(bucket='physical'),(42,D('1425.50')))
        with self.assertRaises(ValidationError): s.reconcile(self.p,self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='waste',rolls=1,yards=D('30')))
        s.reconcile(self.p,self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='waste',rolls=1,yards=D('30'),evidence=self.evidence))
        self.assertEqual(s.balance(bucket='cmt'),(7,D('210')))
    def test_transfer_requires_approved_allocation(self):
        self.received_cmt(); target=self.new_order('PO TRANSFER')
        al=s.save_allocation(self.p,{'order':target,'source_order':self.order,'cmt':self.cmt},[{'lot':self.lot,'warehouse':self.cmt,'rolls':4,'yards':D('120')}])
        data=self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='transfer',rolls=4,yards=D('120'),target_allocation=al.lines.first())
        with self.assertRaises(ValidationError): s.reconcile(self.p,data)
        s.submit_allocation(self.p,al.pk); s.decide_allocation(self.a,al.pk,'approved')
        data['target_allocation']=al.lines.first()
        s.reconcile(self.p,data)
        self.assertEqual(s.balance(bucket='cmt',order=self.order),(6,D('180')))
        self.assertEqual(s.balance(bucket='cmt',order=target),(4,D('120')))
        self.assertEqual(s.balance(bucket='reserved'),(0,D('0')))
    def test_reverse_receipt_rejected_after_reserve(self):
        self.allocate()
        with self.assertRaises(ValidationError): s.reverse(self.p,'receipt',self.receipt.pk,'Salah invoice',self.evidence)
        self.assertEqual(s.balance(bucket='physical'),(50,D('1665.50')))
    def test_reverse_receipt_preserves_history(self):
        self.posted(); s.reverse(self.p,'receipt',self.receipt.pk,'Salah invoice',self.evidence)
        self.assertEqual(s.balance(bucket='physical'),(0,D('0')))
        self.assertEqual(Movement.objects.count(),2)
    def test_reverse_cmt_then_shipment_restores_stock(self):
        sj=self.shipment(); cr=self.received_cmt(sj)
        s.reverse(self.p,'cmt_receipt',cr.pk,'Salah laporan',self.evidence)
        s.reverse(self.p,'shipment',sj.pk,'Salah kirim',self.evidence)
        self.assertEqual(s.balance(bucket='physical'),(50,D('1665.50')))
        self.assertEqual(s.balance(bucket='reserved'),(10,D('300')))
        self.assertEqual(s.balance(bucket='transit'),(0,D('0')))
        self.assertEqual(s.balance(bucket='cmt'),(0,D('0')))
    def test_reverse_warehouse_recalculates_good(self):
        self.allocate(); fs,wr=self.finished(430)
        s.reverse(self.p,'warehouse_receipt',wr.pk,'Salah laporan',self.evidence)
        self.assertEqual(s.totals(self.order)['good'],0)
        fs.refresh_from_db(); self.assertEqual(fs.status,'dispatched')
    def test_adjustment_requires_special_permission_and_free_stock(self):
        self.allocate(40,'1300')
        with self.assertRaises(ValidationError): s.adjust(self.p,self.lot,self.wh,'out',11,D('10'),'Koreksi',self.evidence)
        self.p.can_adjust=False; self.p.save()
        with self.assertRaises(PermissionDenied): s.adjust(self.p,self.lot,self.wh,'in',1,D('10'),'Koreksi',self.evidence)
    def test_closed_order_rejects_progress_and_reopen_audited(self):
        self.received_cmt(); self.finished(1072)
        s.reconcile(self.p,self.report(order=self.order,cmt=self.cmt,lot=self.lot,action='consumed',rolls=10,yards=D('300')))
        s.order_action(self.p,self.order.pk,'close')
        with self.assertRaises(ValidationError): s.save_progress(self.p,self.report(order=self.order,cmt=self.cmt,stage='qc',quantity=1,reject=0))
        s.order_action(self.p,self.order.pk,'reopen','Perbaikan laporan')
        self.assertTrue(Audit.objects.filter(action='reopen').exists())
    def test_float_and_precision_rejected(self):
        for r,y in [(1,0.1),(1.5,'1.00'),(1,'1.001'),(1,'NaN'),(-1,'2'),(1,'1000000000000')]:
            with self.assertRaises(ValidationError): s.qty(r,y)
    def test_controlled_po_number_blocks_separator_duplicates(self):
        for candidate in ['po0111','PO-0111',' PO / 0111 ']:
            with self.assertRaises(ValidationError):
                s.save_order(self.p,{'number':candidate,'product':self.product,'target':10,'cmts':[self.cmt]})
    def test_master_normalization_and_inactive_selection(self):
        with self.assertRaises(ValidationError): s.save_master(self.p,{'kind':'material','code':' inesa ','name':'Duplicate'})
        self.material.active=False; self.material.save()
        self.client.force_login(self.p)
        response=self.client.get('/receipts/new/')
        self.assertEqual(response.status_code,200)
        self.assertNotContains(response,f'<option value="{self.material.pk}">INESA')
    def test_ledger_audit_decisions_append_only_at_database(self):
        self.allocate()
        for model in [Movement,Audit,Decision]:
            obj=model.objects.first()
            with self.assertRaises(ValidationError): obj.delete()
            with self.assertRaises(DatabaseError):
                with transaction.atomic(): model.objects.filter(pk=obj.pk).delete()
    def test_login_required_and_csrf(self):
        client=Client(enforce_csrf_checks=True)
        self.assertEqual(client.get('/orders/').status_code,302)
        client.force_login(self.p)
        response=client.post(f'/receipts/{self.receipt.pk}/action/post/',{'token':str(uuid.uuid4())})
        self.assertEqual(response.status_code,403)
    def test_login_throttle_does_not_lock_other_accounts_behind_same_proxy(self):
        import hashlib
        for key in ['client-account:127.0.0.1:purchasing','account:purchasing']:
            LoginAttempt.objects.create(key=hashlib.sha256(key.encode()).hexdigest(),failures=8)
        response=self.client.post('/login/',{'username':'purchasing','password':'wrong'})
        self.assertEqual(response.status_code,429)
        response=self.client.post('/login/',{'username':'another-account','password':'wrong'})
        self.assertEqual(response.status_code,200)
    def test_all_views_render_for_roles(self):
        sj=self.shipment(); self.received_cmt(sj); fs,wr=self.finished(20)
        for user in [self.p,self.a,self.m,self.admin]:
            self.client.force_login(user)
            for path in ['/','/receipts/','/orders/','/stock/','/movements/','/allocations/','/approvals/','/shipments/','/cmt/','/progress/','/finished/','/warehouse/','/reconciliation/','/exceptions/','/corrections/','/masters/','/reports/',f'/orders/{self.order.pk}/',f'/receipts/{self.receipt.pk}/',f'/allocations/{sj.allocation_id}/',f'/shipments/{sj.pk}/',f'/stock/{self.lot.pk}/',f'/finished/{fs.pk}/',f'/warehouse/{wr.pk}/']:
                with self.subTest(user=user.role,path=path): self.assertEqual(self.client.get(path).status_code,200)
        self.client.force_login(self.p)
        for path in ['/receipts/new/','/orders/new/','/allocations/new/','/masters/new/','/progress/new/','/finished/new/','/warehouse/new/','/reconciliation/new/','/adjustment/',f'/shipments/new/?allocation={sj.allocation_id}',f'/cmt/new/?line={sj.lines.first().pk}']:
            with self.subTest(path=path): self.assertEqual(self.client.get(path).status_code,200)
    def test_receipt_form_post_creates_draft_and_duplicate_retry(self):
        self.client.force_login(self.p)
        data={'vendor':self.vendor.pk,'invoice':'FORM-001','warehouse':self.wh.pk,'received_date':self.today.isoformat(),'invoice_total':'1000.25','token':str(uuid.uuid4()),'version':1,'lines-TOTAL_FORMS':1,'lines-INITIAL_FORMS':0,'lines-MAX_NUM_FORMS':100,'lines-MIN_NUM_FORMS':0,'lines-0-material':self.material.pk,'lines-0-color':self.color.pk,'lines-0-unit':self.unit.pk,'lines-0-warehouse':self.wh.pk,'lines-0-rolls':3,'lines-0-yards':'100.25'}
        response=self.client.post('/receipts/new/',data)
        self.assertEqual(response.status_code,302)
        self.assertEqual(Receipt.objects.get(invoice='FORM-001').lines.count(),1)
        response=self.client.post('/receipts/new/',data)
        self.assertIn(response.status_code,[302,400])
        self.assertEqual(Receipt.objects.filter(invoice='FORM-001').count(),1)
    def test_upload_mime_and_limit(self):
        with self.assertRaises(ValidationError): storage.upload(self.p,SimpleUploadedFile('x.png',b'<script>alert(1)</script>',content_type='image/png'))
        with self.assertRaises(ValidationError): storage.upload(self.p,SimpleUploadedFile('x.exe',b'%PDF-1.7',content_type='application/pdf'))
    def test_account_permissions_and_admin_has_no_operational_access(self):
        with self.assertRaises(PermissionDenied): save_account(self.p,{},None)
        self.client.force_login(self.m); self.assertEqual(self.client.get('/accounts/').status_code,403)
        self.client.force_login(self.admin); self.assertEqual(self.client.get('/accounts/new/').status_code,200)
        with self.assertRaises(PermissionDenied): s.save_master(self.admin,{'kind':'vendor','code':'NO','name':'NO'})
    def test_stock_po_filter_does_not_multiply_shared_lot(self):
        self.allocate(5,'150'); self.allocate(5,'150')
        from .reports import dataset
        headers,rows=dataset('stock',{'order':self.order},self.p)
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['cells'][3]['value'],'1.665,50')
        self.assertEqual(rows[0]['cells'][4]['value'],'300,00')
        headers,rows=dataset('stock',{'status':'physical'},self.p)
        self.assertEqual(rows[0]['cells'][4]['value'],'300,00')
    def test_card_balance_preserves_receipt_when_filtering_po(self):
        self.shipment()
        from .reports import dataset
        headers,rows=dataset('movements',{'order':self.order},self.p)
        physical=[r for r in rows if r['object'].bucket=='physical']
        self.assertEqual(physical[0]['cells'][7]['value'],'1.365,50')
    def test_form_pagination_export_and_retry(self):
        self.client.force_login(self.p)
        for index in range(27):
            s.save_master(self.p,{'kind':'vendor','code':f'P-{index:03d}','name':f'Vendor {index}'})
        response=self.client.get('/masters/',{'kind':'vendor','page':2,'q':'Vendor'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.context['count'],28)
        self.assertEqual(len(response.context['page'].object_list),3)
        response=self.client.get('/masters/export/',{'kind':'vendor','page':2,'q':'Vendor','_page':1})
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(wb.active.max_row,29)

@skipUnless(connection.vendor=='postgresql','Concurrency membutuhkan PostgreSQL asli.')
class ConcurrencyTests(Fixture,TransactionTestCase):
    reset_sequences=True
    def setUp(self): self.build()
    def test_at04_competing_approvals_cannot_overreserve(self):
        a1=self.allocate(30,'1000',approve=False)
        a2=self.allocate(30,'1000',order=self.new_order('PO RACE'),approve=False)
        barrier=threading.Barrier(2)
        def run(pk):
            try:
                actor=User.objects.get(pk=self.a.pk); barrier.wait(timeout=10)
                s.decide_allocation(actor,pk,'approved'); return 'ok'
            except ValidationError: return 'insufficient'
            finally: connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            result=list(pool.map(run,[a1.pk,a2.pk]))
        self.assertCountEqual(result,['ok','insufficient'])
        self.assertEqual(s.balance(bucket='reserved'),(30,D('1000')))
    def test_simultaneous_duplicate_post_is_idempotent(self):
        token=uuid.uuid4(); barrier=threading.Barrier(2)
        def run(_):
            try:
                actor=User.objects.get(pk=self.p.pk); barrier.wait(timeout=10)
                return s.post_receipt(actor,self.receipt.pk,key=token).pk
            finally: connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool: result=list(pool.map(run,[1,2]))
        self.assertEqual(result[0],result[1]); self.assertEqual(Movement.objects.count(),1)
