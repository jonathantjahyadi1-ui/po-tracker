from collections import defaultdict
from datetime import date
from functools import wraps
from io import BytesIO
from urllib.parse import urlencode
from uuid import uuid4

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db import DatabaseError, OperationalError, transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import content_disposition_header
from django.views.decorators.http import require_GET, require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .forms import PaymentForm
from .models import Invoice, InvoiceAttachment, InvoicePayment
from .payment_services import (
    payment_queryset,
    payment_summary,
    record_payment,
    require_director,
)

STATUSES = [
    ('unpaid', 'Unpaid'),
    ('dicicil', 'Dicicil'),
    ('paid', 'Paid'),
    ('reconciliation', 'Perlu rekonsiliasi'),
]


def payment_access(view):
    @login_required
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not request.user.is_active:
            raise PermissionDenied
        return view(request, *args, **kwargs)

    return wrapped


def read_filters(request):
    filters = {
        'q': request.GET.get('q', '').strip(),
        'status': request.GET.get('status', ''),
        'date_from': None,
        'date_to': None,
    }
    if filters['status'] and filters['status'] not in dict(STATUSES):
        raise ValidationError('Status Payment tidak dikenal. Pilih status yang tersedia.')
    for key in ('date_from', 'date_to'):
        raw = request.GET.get(key, '')
        if raw:
            try:
                filters[key] = date.fromisoformat(raw)
            except ValueError:
                raise ValidationError('Periode Invoice harus menggunakan tanggal yang valid.')
    if filters['date_from'] and filters['date_to'] and filters['date_from'] > filters['date_to']:
        raise ValidationError('Awal Periode Invoice tidak boleh melewati tanggal akhirnya.')
    return filters


def filter_query(filters):
    return urlencode({key: str(value) for key, value in filters.items() if value})


def can_pay(user):
    return user.role == 'direktur' and not user.is_superuser and user.is_active


@payment_access
@require_GET
def payment_list(request):
    error = None
    try:
        filters = read_filters(request)
        queryset = payment_queryset(filters)
    except ValidationError as exc:
        filters = {'q': request.GET.get('q', ''), 'status': '', 'date_from': None, 'date_to': None}
        queryset = Invoice.objects.none()
        error = '; '.join(exc.messages)
    page = Paginator(queryset, 50).get_page(request.GET.get('page'))
    for invoice in page.object_list:
        invoice.payment = payment_summary(invoice)
    return render(
        request,
        'payment_list.html',
        {
            'title': 'Payment',
            'active': 'payment',
            'page_obj': page,
            'filters': filters,
            'statuses': STATUSES,
            'filter_query': filter_query(filters),
            'filter_error': error,
            'can_pay': can_pay(request.user),
        },
        status=400 if error else 200,
    )


@payment_access
@require_GET
def payment_detail(request, pk):
    invoice = get_object_or_404(Invoice.objects.select_related('vendor', 'dibuat_oleh'), pk=pk)
    payments = list(invoice.payments.select_related('recorded_by').defer('proof_content'))
    summary = payment_summary(invoice, payments)
    attachment = InvoiceAttachment.objects.filter(invoice=invoice).defer('content').first()
    return render(
        request,
        'payment_detail.html',
        {
            'title': f'Payment · {invoice.nomor}',
            'active': 'payment',
            'invoice': invoice,
            'summary': summary,
            'payments': payments,
            'attachment': attachment,
            'can_pay': can_pay(request.user)
            and not invoice.dibatalkan
            and summary.status_key in ('unpaid', 'dicicil'),
        },
    )


@payment_access
@require_http_methods(['GET', 'POST'])
def payment_create(request, pk):
    require_director(request.user)
    invoice = get_object_or_404(Invoice.objects.select_related('vendor'), pk=pk)
    summary = payment_summary(invoice)
    if request.method == 'GET' and (
        invoice.dibatalkan or summary.status_key not in ('unpaid', 'dicicil')
    ):
        messages.error(request, 'Invoice ini belum dapat menerima pembayaran atau sudah Paid.')
        return redirect('payment_detail', pk=pk)
    form = PaymentForm(
        request.POST if request.method == 'POST' else None,
        request.FILES if request.method == 'POST' else None,
        initial={
            'kind': 'lunas',
            'amount': summary.remaining,
            'remaining_seen': summary.remaining,
            'request_id': uuid4().hex,
        },
    )
    status = 200
    if request.method == 'POST':
        status = 400
        if form.is_valid():
            try:
                record_payment(invoice, form.cleaned_data, request.user)
            except ValidationError as exc:
                form.add_error(None, exc)
                status = 409 if any(error.code == 'conflict' for error in exc.error_list) else 400
            except OperationalError:
                form.add_error(
                    None, 'Pembayaran lain sedang diproses. Tinjau sisa terbaru dan coba lagi.'
                )
                status = 409
            except (DatabaseError, OSError):
                form.add_error(
                    None, 'Pembayaran belum tersimpan. Periksa isian dan unggah ulang bukti.'
                )
                status = 503
            else:
                messages.success(
                    request, 'Pembayaran tersimpan. Posisi tagihan dan riwayat diperbarui.'
                )
                return redirect('payment_detail', pk=pk)
        invoice.refresh_from_db()
        summary = payment_summary(invoice)
    return render(
        request,
        'payment_form.html',
        {
            'title': f'Bayar · {invoice.nomor}',
            'active': 'payment',
            'invoice': invoice,
            'summary': summary,
            'form': form,
        },
        status=status,
    )


@payment_access
@require_GET
def payment_proof(request, pk):
    payment = get_object_or_404(InvoicePayment, pk=pk)
    response = HttpResponse(payment.proof_content, content_type=payment.proof_content_type)
    response['Content-Disposition'] = content_disposition_header(
        request.GET.get('download') == '1', payment.proof_filename
    )
    response['Content-Length'] = payment.proof_size
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


def append_record(sheet, values):
    sheet.append(values)
    # User-entered identifiers remain text, including strings starting with '='.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def style_sheet(sheet, money_columns, date_columns):
    sheet.freeze_panes = 'A6'
    sheet.auto_filter.ref = f'A5:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}'
    for cell in sheet[5]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1C6B50')
    for row in sheet.iter_rows(min_row=6):
        for column in money_columns:
            row[column - 1].number_format = '#,##0.00'
        for column in date_columns:
            row[column - 1].number_format = 'dd/mm/yyyy'
    for column in sheet.columns:
        width = max(len(str(cell.value or '')) for cell in column[4:]) + 2
        sheet.column_dimensions[column[0].column_letter].width = min(48, max(16, width))


@payment_access
@require_GET
def payment_export(request):
    try:
        filters = read_filters(request)
    except ValidationError as exc:
        messages.error(request, '; '.join(exc.messages))
        return redirect('payment_list')
    # Take parent locks before collecting history so both sheets use the same position.
    # Writers lock these same invoice rows before changing totals or adding payments.
    with transaction.atomic():
        selected_ids = list(payment_queryset(filters).values_list('pk', flat=True))
        invoices = list(
            Invoice.objects.filter(pk__in=selected_ids)
            .order_by('pk')
            .select_for_update(of=('self',))
            .select_related('vendor')
        )
        payments = list(
            InvoicePayment.objects.filter(invoice_id__in=selected_ids)
            .select_related('recorded_by')
            .defer('proof_content')
            .order_by('invoice_id', 'payment_date', 'created_at', 'pk')
        )
        exported_at = timezone.localtime()
        grouped = defaultdict(list)
        for payment in payments:
            grouped[payment.invoice_id].append(payment)
        for invoice in invoices:
            invoice.payment = payment_summary(invoice, grouped[invoice.pk])
    workbook = Workbook()
    recap = workbook.active
    recap.title = 'Rekap Invoice'
    history = workbook.create_sheet('Riwayat Pembayaran')
    filter_text = (
        f'Pencarian: {filters["q"] or "Semua"}; '
        f'Status: {dict(STATUSES).get(filters["status"], "Semua")}; '
        f'Periode Invoice: {filters["date_from"] or "Semua"} s.d. '
        f'{filters["date_to"] or "Semua"}'
    )
    for sheet in (recap, history):
        append_record(sheet, ['Payment Invoice — PO Tracker'])
        append_record(
            sheet, ['Waktu ekspor (Asia/Jakarta)', exported_at.strftime('%d/%m/%Y %H:%M:%S')]
        )
        append_record(sheet, ['Filter', filter_text])
        append_record(sheet, ['Jumlah invoice', len(invoices)])
    append_record(
        recap,
        [
            'ID invoice',
            'Nomor invoice',
            'Vendor',
            'Tanggal invoice',
            'Total tagihan',
            'Total dibayar',
            'Sisa tagihan',
            'Progres (%)',
            'Status',
            'Tanggal pembayaran terakhir',
        ],
    )
    append_record(
        history,
        [
            'ID transaksi',
            'ID invoice',
            'Nomor invoice',
            'Vendor',
            'Tanggal pembayaran',
            'Waktu dicatat (Asia/Jakarta)',
            'Jenis',
            'Nominal transaksi',
            'Nama pencatat',
            'Tautan bukti',
        ],
    )
    for invoice in invoices:
        summary = invoice.payment
        append_record(
            recap,
            [
                invoice.pk,
                invoice.nomor,
                invoice.vendor.name,
                invoice.tanggal,
                invoice.total_rp,
                summary.paid,
                summary.remaining,
                summary.progress,
                summary.status,
                summary.last_payment,
            ],
        )
        for payment in grouped[invoice.pk]:
            proof_url = request.build_absolute_uri(reverse('payment_proof', args=[payment.pk]))
            append_record(
                history,
                [
                    payment.pk,
                    invoice.pk,
                    invoice.nomor,
                    invoice.vendor.name,
                    payment.payment_date,
                    timezone.localtime(payment.created_at).replace(tzinfo=None),
                    payment.get_kind_display(),
                    payment.amount,
                    payment.recorded_by.get_full_name() or payment.recorded_by.username,
                    proof_url,
                ],
            )
            cell = history.cell(history.max_row, 10)
            cell.hyperlink = proof_url
            cell.style = 'Hyperlink'
    style_sheet(recap, (5, 6, 7, 8), (4, 10))
    style_sheet(history, (8,), (5,))
    for row in history.iter_rows(min_row=6):
        row[5].number_format = 'dd/mm/yyyy hh:mm:ss'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    name = exported_at.strftime('payment-%Y%m%d-%H%M%S.xlsx')
    response['Content-Disposition'] = content_disposition_header(True, name)
    response['Cache-Control'] = 'private, no-store'
    return response
