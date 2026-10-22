from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
from unittest.mock import patch
from uuid import uuid4

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import DatabaseError, close_old_connections, connection
from django.db.models import Sum
from django.test import TestCase, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .defect_services import (
    cancel_defect,
    cancel_refund,
    record_defect,
    record_refund,
    return_defect_item,
)
from .models import (
    Alokasi,
    FabricDefect,
    FabricDefectItem,
    Hasil,
    Invoice,
    Log,
    Master,
    Po,
    Roll,
    User,
    VendorRefund,
)
from .payment_services import payment_queryset, payment_summary, record_payment
from .services import (
    ajukan_alokasi,
    batalkan_invoice,
    po_status,
    putuskan_alokasi,
    simpan_invoice,
    tautkan_po,
    terima_alokasi,
)
from .test_payments import payment_data
from .test_size_helpers import assign_test_po, standard_product


class DefectTests(TestCase):
    def setUp(self):
        self.buyer = User.objects.create_user('bs-buyer', role='purchasing')
        self.director = User.objects.create_user('bs-director', role='direktur')
        self.admin = User.objects.create_user('bs-admin', role='admin')
        self.accounting = User.objects.create_user('bs-reader', role='accounting')
        self.today = timezone.localdate()
        self.cmt = Master.objects.create(kind='cmt', name='CMT BS')
        self.invoice, self.roll = self.source('INV-BS')

    def source(self, number, vendor='Vendor BS', total='100000000'):
        invoice = simpan_invoice(
            {
                'vendor': vendor,
                'nomor': number,
                'tanggal': self.today,
                'total_rp': total,
            },
            [{'material': 'Cotton BS', 'rows': [{'color': 'Cream BS', 'yards': [Decimal(100)]}]}],
            self.buyer,
        )
        roll = invoice.roll_set.select_related('material', 'color').get()
        roll.material.price_per_yard = Decimal('1000000')
        roll.material.save(update_fields=['price_per_yard'])
        allocation = ajukan_alokasi([roll.pk], self.cmt, self.buyer)
        putuskan_alokasi(allocation, 'acc', self.buyer, tgl_kirim=self.today)
        terima_alokasi(allocation, [roll.pk], self.today, self.buyer)
        roll.refresh_from_db()
        return invoice, roll

    def bs(self, yard='10', **changes):
        return record_defect(
            **{
                'cmt': self.cmt,
                'rows': [{'roll': self.roll.pk, 'yard': yard}],
                'date': self.today,
                'note': 'Cacat kain',
                'request_id': uuid4().hex,
                'user': self.buyer,
                **changes,
            }
        )

    def pay(self, amount, invoice=None):
        invoice = invoice or self.invoice
        return record_payment(
            invoice,
            payment_data(
                amount=str(amount),
                seen=str(invoice.total_rp),
            ),
            self.director,
        )

    def refund(self, amount, **changes):
        return record_refund(
            **{
                'invoice': self.invoice,
                'amount': str(amount),
                'date': self.today,
                'note': 'Transfer vendor',
                'request_id': uuid4().hex,
                'user': self.director,
                **changes,
            }
        )

    def test_before_po_partial_and_full_stock_preserve_original_receipt(self):
        first = self.bs('15')
        self.roll.refresh_from_db()
        item = first.items.get()
        self.assertIsNone(item.po_id)
        self.assertEqual(
            (item.invoice_id, item.vendor_id), (self.invoice.pk, self.invoice.vendor_id)
        )
        self.assertEqual((self.roll.yard, self.roll.bs_yard, self.roll.usable_yard), (100, 15, 85))
        self.assertEqual(self.roll.bs_label, 'Sebagian BS')
        self.bs('85')
        self.roll.refresh_from_db()
        self.assertEqual((self.roll.usable_yard, self.roll.bs_label), (0, 'BS'))

    def test_decimal_comma_and_price_snapshot(self):
        material = self.roll.material
        material.price_per_yard = Decimal('20000')
        material.save(update_fields=['price_per_yard'])
        old = self.bs('12,5').items.get()
        self.assertEqual((old.yard, old.price, old.amount), (Decimal('12.5'), 20000, 250000))
        material.price_per_yard = Decimal('40000')
        material.save(update_fields=['price_per_yard'])
        new = self.bs('1').items.get()
        old.refresh_from_db()
        self.assertEqual(
            (old.price, old.amount, new.price, new.amount), (20000, 250000, 40000, 40000)
        )

    def test_money_rounds_each_item_half_up_and_never_rounds_yards_to_pcs(self):
        self.roll.material.price_per_yard = Decimal('0.01')
        self.roll.material.save(update_fields=['price_per_yard'])
        item = self.bs('1,5').items.get()
        self.assertEqual((item.yard, item.amount), (Decimal('1.5'), Decimal('0.02')))

    def test_invalid_yards_missing_price_or_excess_value_are_atomic(self):
        before = Log.objects.count()
        for value in ('0', '-1', 'NaN', '101', '1.1234567', 'abc'):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                self.bs(value)
        self.roll.material.price_per_yard = None
        self.roll.material.save(update_fields=['price_per_yard'])
        with self.assertRaisesMessage(ValidationError, 'Harga per Yard'):
            self.bs()
        self.roll.material.price_per_yard = Decimal('2000000')
        self.roll.material.save(update_fields=['price_per_yard'])
        with self.assertRaisesMessage(ValidationError, 'melampaui tagihan'):
            self.bs('60')
        self.assertEqual(Log.objects.count(), before)
        self.assertFalse(FabricDefect.objects.exists())
        self.roll.refresh_from_db()
        self.assertEqual(self.roll.bs_yard, 0)

    def test_all_prd_invoice_positions_keep_cash_history(self):
        for index, paid, remaining, refund in (
            (0, 0, 90000000, 0),
            (1, 60000000, 30000000, 0),
            (2, 100000000, 0, 10000000),
            (3, 95000000, 0, 5000000),
        ):
            invoice, roll = self.source(f'POSITION-{index}')
            if paid:
                self.pay(paid, invoice)
            self.bs(rows=[{'roll': roll.pk, 'yard': '10'}])
            summary = payment_summary(invoice)
            self.assertEqual(
                (
                    summary.net_total,
                    summary.paid,
                    summary.remaining,
                    summary.refund_due,
                    summary.refund_received,
                ),
                (90000000, paid, remaining, refund, 0),
            )
            self.assertEqual(invoice.payments.aggregate(total=Sum('amount'))['total'] or 0, paid)
            if not remaining:
                self.assertEqual(summary.status, 'Paid')
            self.assertEqual(invoice.total_rp, 100000000)

    def test_payment_after_bs_cannot_exceed_net_remaining(self):
        self.pay(60000000)
        self.bs()
        with self.assertRaisesMessage(ValidationError, 'melebihi sisa tagihan'):
            self.pay(40000000)
        record_payment(
            self.invoice,
            payment_data(
                amount='30000000',
                kind='lunas',
                seen='30000000',
            ),
            self.director,
        )
        self.assertEqual(payment_summary(self.invoice).status, 'Paid')
        self.assertEqual(self.invoice.payments.count(), 2)

    def test_partial_refunds_additional_bs_and_next_invoice_are_independent(self):
        self.pay(100000000)
        defect = self.bs()
        first = self.refund(7000000)
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.refund_remaining, summary.refund_status), (3000000, 'Refund diterima sebagian')
        )
        self.bs('5')
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.refund_due, summary.refund_received, summary.refund_remaining),
            (15000000, 7000000, 8000000),
        )
        self.assertIsNone(defect.items.get().returned_at)
        self.refund(8000000)
        summary = payment_summary(self.invoice)
        self.assertEqual(
            (summary.paid, summary.status, summary.refund_remaining, summary.refund_status),
            (100000000, 'Paid', 0, 'Refund diterima'),
        )
        self.assertEqual(first.amount, 7000000)
        next_invoice, _ = self.source('NEXT-INVOICE', total='60000000')
        self.assertEqual(
            (payment_summary(next_invoice).net_total, payment_summary(next_invoice).remaining),
            (60000000, 60000000),
        )

    def test_next_invoice_before_refund_also_stands_alone(self):
        self.pay(100000000)
        self.bs()
        invoice, _ = self.source('NEXT-BEFORE-REFUND', total='60000000')
        self.assertEqual(payment_summary(invoice).remaining, 60000000)

    def test_return_does_not_subtract_usable_yard_twice_or_receive_money(self):
        self.pay(100000000)
        item = self.bs('15').items.get()
        return_defect_item(item, self.today, 'SJ RETUR', self.buyer)
        return_defect_item(item, self.today, 'SJ RETUR', self.buyer)
        self.roll.refresh_from_db()
        self.assertEqual((self.roll.usable_yard, self.roll.status), (85, 'diterima'))
        self.assertEqual(payment_summary(self.invoice).refund_due, 15000000)
        self.assertFalse(VendorRefund.objects.exists())
        self.assertEqual(Log.objects.filter(aksi='retur BS').count(), 1)

    def test_full_return_has_explicit_roll_status_and_is_not_allocatable(self):
        item = self.bs('100').items.get()
        return_defect_item(item, self.today, '', self.buyer)
        self.roll.refresh_from_db()
        self.assertEqual(
            (self.roll.status, self.roll.yard, self.roll.usable_yard), ('diretur', 100, 0)
        )
        with self.assertRaises(ValidationError):
            ajukan_alokasi([self.roll.pk], self.cmt, self.buyer)

    def test_multi_vendor_items_use_origin_invoice_and_return_individually(self):
        other, roll = self.source('INV-OTHER-VENDOR', vendor='Vendor Lain')
        self.pay(100000000, other)
        defect = self.bs(
            rows=[{'roll': self.roll.pk, 'yard': '15'}, {'roll': roll.pk, 'yard': '10'}]
        )
        first, second = list(defect.items.order_by('roll_id'))
        return_defect_item(first, self.today, '', self.buyer)
        second.refresh_from_db()
        self.assertIsNone(second.returned_at)
        self.assertEqual(
            (first.vendor_id, second.vendor_id), (self.invoice.vendor_id, other.vendor_id)
        )
        self.assertEqual(
            (payment_summary(self.invoice).remaining, payment_summary(other).refund_due),
            (85000000, 10000000),
        )

    def test_bs_and_refund_replay_reject_changed_payload_and_do_not_duplicate(self):
        self.pay(100000000)
        key = uuid4().hex
        first = self.bs(request_id=key)
        self.assertEqual(self.bs(request_id=key).pk, first.pk)
        with self.assertRaisesMessage(ValidationError, 'data berbeda'):
            self.bs('11', request_id=key)
        key = uuid4().hex
        received = self.refund(7000000, request_id=key)
        self.assertEqual(self.refund(7000000, request_id=key).pk, received.pk)
        with self.assertRaisesMessage(ValidationError, 'data berbeda'):
            self.refund(6000000, request_id=key)
        for _ in range(3):
            self.assertEqual(payment_summary(self.invoice).refund_remaining, 3000000)
        self.assertEqual(
            (
                FabricDefect.objects.count(),
                FabricDefectItem.objects.count(),
                VendorRefund.objects.count(),
            ),
            (1, 1, 1),
        )

    def test_refund_limits_and_permissions_are_checked_in_service(self):
        self.pay(100000000)
        self.bs()
        for amount in ('0', '10000001', '-1'):
            with self.subTest(amount=amount), self.assertRaises(ValidationError):
                self.refund(amount)
        for user in (self.buyer, self.accounting):
            with self.subTest(user=user.role), self.assertRaises(PermissionDenied):
                self.refund(1, user=user)
        self.refund(10000000, user=self.admin)
        with self.assertRaises(ValidationError):
            self.refund(1)

    def test_wrong_cmt_unreceived_and_used_production_are_rejected(self):
        other_cmt = Master.objects.create(kind='cmt', name='CMT Lain')
        with self.assertRaisesMessage(ValidationError, 'CMT sumber'):
            self.bs(cmt=other_cmt)
        Roll.objects.filter(pk=self.roll.pk).update(status='dikirim')
        with self.assertRaises(ValidationError):
            self.bs()
        Roll.objects.filter(pk=self.roll.pk).update(status='diterima')
        allocation = Alokasi.objects.get(pk=self.roll.alokasi_id)
        po = assign_test_po(allocation, 'PO USED BS', self.buyer)
        Hasil.objects.create(po=po, material=self.roll.material, color=self.roll.color, pcs=10)
        with self.assertRaisesMessage(ValidationError, 'hasil produksi'):
            self.bs()
        self.assertFalse(FabricDefect.objects.exists())

    def test_po_estimates_exclude_bs_without_changing_pcs_rounding(self):
        self.bs('15')
        allocation = Alokasi.objects.get(pk=self.roll.alokasi_id)
        product = standard_product(self.buyer, usage='2')
        po = tautkan_po(allocation, 'PO GOOD YARD', self.buyer, product=product)
        summary = po_status(po)
        self.assertEqual(
            (summary['rows'][0]['allocated_yard'], summary['rows'][0]['estimate']), (85, 43)
        )
        extra = self.bs('5').items.get()
        self.assertEqual(extra.po_id, po.pk)
        self.assertEqual(po_status(po)['rows'][0]['estimate'], 40)

    def test_whole_bs_cannot_be_assigned_to_po(self):
        self.bs('100')
        with self.assertRaises(ValidationError):
            tautkan_po(
                Alokasi.objects.get(pk=self.roll.alokasi_id),
                'PO ALL BS',
                self.buyer,
                product=Master.objects.create(kind='product', name='Produk BS'),
            )

    def test_cancel_reverses_stock_finance_and_preserves_audit(self):
        self.pay(60000000)
        defect = self.bs('15')
        cancel_defect(defect, 'Salah yard', self.admin)
        cancel_defect(defect, 'Salah yard', self.admin)
        self.roll.refresh_from_db()
        self.assertEqual(
            (self.roll.usable_yard, payment_summary(self.invoice).remaining), (100, 40000000)
        )
        self.assertEqual((FabricDefect.objects.count(), FabricDefectItem.objects.count()), (1, 1))
        self.assertEqual(Log.objects.filter(aksi='batalkan BS').count(), 1)

    def test_cancel_blocks_returned_stock_and_refund_below_receipts(self):
        self.pay(100000000)
        defect = self.bs()
        refund = self.refund(7000000)
        with self.assertRaisesMessage(ValidationError, 'Refund sudah diterima'):
            cancel_defect(defect, 'Koreksi', self.admin)
        self.roll.refresh_from_db()
        self.assertEqual(self.roll.bs_yard, 10)
        cancel_refund(refund, 'Salah penerimaan', self.admin)
        self.assertEqual(payment_summary(self.invoice).refund_received, 0)
        return_defect_item(defect.items.get(), self.today, '', self.buyer)
        with self.assertRaisesMessage(ValidationError, 'sudah diretur'):
            cancel_defect(defect, 'Koreksi', self.admin)

    def test_cancelling_additional_bs_can_preserve_valid_received_refund(self):
        self.pay(100000000)
        self.bs()
        extra = self.bs('5')
        self.refund(7000000)
        cancel_defect(extra, 'Duplikat laporan lapangan', self.admin)
        self.assertEqual(payment_summary(self.invoice).refund_remaining, 3000000)

    def test_failed_log_rolls_back_bs_and_refund_entirely(self):
        with patch('tracking.defect_services.log', side_effect=DatabaseError('failure')):
            with self.assertRaises(DatabaseError):
                self.bs()
        self.roll.refresh_from_db()
        self.assertEqual(self.roll.bs_yard, 0)
        self.assertFalse(FabricDefect.objects.exists())
        self.pay(100000000)
        self.bs()
        with patch('tracking.defect_services.log', side_effect=DatabaseError('failure')):
            with self.assertRaises(DatabaseError):
                self.refund(10000000)
        self.assertFalse(VendorRefund.objects.exists())

    def test_invalid_second_invoice_roll_rolls_back_the_whole_multi_roll_request(self):
        other, roll = self.source('ATOMIC-OTHER')
        roll.material.price_per_yard = None
        roll.material.save(update_fields=['price_per_yard'])
        with self.assertRaises(ValidationError):
            self.bs(rows=[{'roll': self.roll.pk, 'yard': '10'}, {'roll': roll.pk, 'yard': '10'}])
        self.assertFalse(FabricDefect.objects.exists())
        self.roll.refresh_from_db()
        roll.refresh_from_db()
        self.assertEqual((self.roll.bs_yard, roll.bs_yard), (0, 0))
        self.assertEqual(payment_summary(other).remaining, other.total_rp)

    def test_repeated_roll_in_one_request_and_overlapping_existing_bs_are_rejected(self):
        with self.assertRaisesMessage(ValidationError, 'dua kali'):
            self.bs(
                rows=[{'roll': self.roll.pk, 'yard': '10'}, {'roll': self.roll.pk, 'yard': '10'}]
            )
        self.bs('80')
        with self.assertRaisesMessage(ValidationError, 'melebihi sisa'):
            self.bs('30')
        self.roll.refresh_from_db()
        self.assertEqual(self.roll.bs_yard, 80)

    def test_existing_bs_locks_original_invoice_identity_and_value(self):
        self.bs()
        with self.assertRaises(ValidationError):
            simpan_invoice(
                {
                    'vendor': self.invoice.vendor.name,
                    'nomor': self.invoice.nomor,
                    'tanggal': self.today,
                    'total_rp': '90000000',
                },
                [],
                self.buyer,
                self.invoice,
            )
        with self.assertRaises(ValidationError):
            batalkan_invoice(self.invoice, self.buyer)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.total_rp, 100000000)

    def test_direct_http_permissions_and_post_only_mutations(self):
        self.pay(100000000)
        defect = self.bs()
        self.client.force_login(self.buyer)
        refund_url = reverse('refund_create', args=[self.invoice.pk])
        self.assertEqual(self.client.post(refund_url, {}).status_code, 403)
        self.assertEqual(
            self.client.post(
                reverse('defect_cancel', args=[defect.pk]), {'reason': 'Unauthorized'}
            ).status_code,
            403,
        )
        for name, pk in (('defect_cancel', defect.pk), ('defect_return', defect.items.get().pk)):
            self.assertEqual(self.client.get(reverse(name, args=[pk])).status_code, 405)
        for user in (self.director, self.accounting):
            self.client.force_login(user)
            self.assertEqual(
                self.client.post(reverse('defect_create', args=[self.cmt.pk]), {}).status_code, 403
            )
            if user == self.accounting:
                self.assertEqual(self.client.post(refund_url, {}).status_code, 403)

    def test_http_bs_error_preserves_input_then_success_and_refund_receipt(self):
        self.pay(100000000)
        self.client.force_login(self.buyer)
        url = reverse('defect_create', args=[self.cmt.pk])
        data = {
            'roll': [str(self.roll.pk)],
            f'yard-{self.roll.pk}': '101',
            'date': self.today.isoformat(),
            'note': 'Jangan hilangkan isian',
            'request_id': uuid4().hex,
        }
        response = self.client.post(url, data)
        self.assertContains(response, 'Jangan hilangkan isian', status_code=400)
        self.assertContains(response, 'value="101"', status_code=400)
        data[f'yard-{self.roll.pk}'] = '10'
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.post(url, data).status_code, 302)
        self.assertEqual(self.client.get(response.url).status_code, 200)
        self.client.force_login(self.director)
        response = self.client.post(
            reverse('refund_create', args=[self.invoice.pk]),
            {
                'amount': '7.000.000',
                'date': self.today.isoformat(),
                'note': 'Transfer',
                'request_id': uuid4().hex,
            },
        )
        self.assertEqual(response.status_code, 302)
        response = self.client.get(reverse('payment_detail', args=[self.invoice.pk]))
        self.assertContains(response, 'Refund diterima sebagian')
        self.assertContains(response, '3.000.000,00')

    def test_payment_filters_badges_views_and_exports_consistent_and_read_only(self):
        self.pay(100000000)
        defect = self.bs()
        self.refund(7000000)
        self.client.force_login(self.buyer)
        before = (
            Log.objects.count(),
            FabricDefectItem.objects.count(),
            VendorRefund.objects.count(),
        )
        for name, pk in (
            ('cmt_detail', self.cmt.pk),
            ('defect_detail', defect.pk),
            ('invoice_detail', self.invoice.pk),
            ('payment_detail', self.invoice.pk),
            ('vendor_detail', self.invoice.vendor_id),
        ):
            response = self.client.get(reverse(name, args=[pk]))
            self.assertEqual(response.status_code, 200, name)
            if name in ('vendor_detail', 'payment_detail', 'invoice_detail'):
                self.assertContains(response, 'Refund diterima sebagian')
        response = self.client.get(reverse('payment_list'), {'status': 'paid'})
        self.assertContains(response, 'Refund diterima sebagian')
        filters = {'q': '', 'status': 'paid', 'date_from': None, 'date_to': None}
        self.assertEqual(payment_queryset(filters).count(), 1)
        filters['status'] = 'reconciliation'
        self.assertEqual(payment_queryset(filters).count(), 0)
        for name, args in (('payment_export', []), ('invoice_export', [self.invoice.pk])):
            response = self.client.get(reverse(name, args=args))
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(response.content))
            position = book['Posisi BS dan Refund']
            self.assertEqual(
                tuple(position.cell(2, col).value for col in range(3, 11)),
                (100000000, 10000000, 90000000, 100000000, 0, 10000000, 7000000, 3000000),
            )
            self.assertEqual(book['Rincian BS']['K2'].value, 10)
            self.assertEqual(book['Penerimaan Refund']['E2'].value, 7000000)
            book.close()
        self.assertEqual(
            before,
            (Log.objects.count(), FabricDefectItem.objects.count(), VendorRefund.objects.count()),
        )

    def test_billion_invoice_and_full_offset_have_no_fictitious_payment(self):
        invoice, roll = self.source('BILLION-BS', total='1000000000')
        roll.material.price_per_yard = Decimal('10000000')
        roll.material.save(update_fields=['price_per_yard'])
        self.pay(1000000000, invoice)
        self.bs(rows=[{'roll': roll.pk, 'yard': '10'}])
        summary = payment_summary(invoice)
        self.assertEqual((summary.net_total, summary.refund_due), (900000000, 100000000))
        self.bs('10')  # The same master price now makes this invoice fully offset.
        summary = payment_summary(self.invoice)
        self.assertEqual((summary.status, summary.remaining, summary.paid), ('Paid', 0, 0))
        self.assertFalse(self.invoice.payments.exists())

    def test_master_price_and_existing_usage_remain_available(self):
        self.client.force_login(self.buyer)
        response = self.client.post(
            reverse('master') + '?tab=material',
            {
                'kind': 'material',
                'id': self.roll.material_id,
                'name': self.roll.material.name,
                'price_per_yard': '20.000',
            },
        )
        self.assertEqual(response.status_code, 302)
        self.roll.material.refresh_from_db()
        self.assertEqual(self.roll.material.price_per_yard, 20000)
        response = self.client.get(reverse('master') + '?tab=material')
        self.assertContains(response, 'Harga per Yard (Rp)')
        self.assertContains(response, 'Pemakaian Paten Bahan')

    def test_no_financial_or_stock_changes_without_purchasing_or_admin_authority(self):
        for user in (self.director, self.admin, self.accounting):
            with self.assertRaises(PermissionDenied):
                self.bs(user=user)
        defect = self.bs()
        for user in (self.buyer, self.director, self.accounting):
            with self.assertRaises(PermissionDenied):
                cancel_defect(defect, 'Forbidden', user)

    def test_vendor_list_uses_aggregate_instead_of_loading_all_rolls(self):
        self.bs()
        self.client.force_login(self.buyer)
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(reverse('vendor_detail', args=[self.invoice.vendor_id]))
        self.assertEqual(response.status_code, 200)
        self.assertLessEqual(len(queries), 6)
        self.assertFalse(
            any(
                '"tracking_roll"."id"' in query['sql'] and 'SELECT "tracking_roll"' in query['sql']
                for query in queries
            )
        )

    def test_default_po_page_only_resolves_the_visible_thirty_orders(self):
        from .views import balance_flags

        Po.objects.bulk_create([Po(nomor=f'PO PAGED {index}', cmt=self.cmt) for index in range(80)])
        self.client.force_login(self.buyer)
        with patch('tracking.views.balance_flags', wraps=balance_flags) as resolver:
            response = self.client.get(reverse('po_cmt_list', args=[self.cmt.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['page_obj'].paginator.count, 80)
        self.assertEqual(len(resolver.call_args.args[0]), 30)


@skipUnless(connection.vendor == 'postgresql', 'Penguncian baris memerlukan PostgreSQL tes.')
class DefectConcurrencyTests(TransactionTestCase):
    def setUp(self):
        self.buyer = User.objects.create_user('race-bs-buyer', role='purchasing')
        self.director = User.objects.create_user('race-bs-director', role='direktur')
        self.cmt = Master.objects.create(kind='cmt', name='Race CMT')
        self.today = timezone.localdate()
        self.vendor = Master.objects.create(kind='vendor', name='Race Vendor')
        self.material = Master.objects.create(
            kind='material', name='Race Fabric', price_per_yard=100
        )
        color = Master.objects.create(kind='color', name='Race Color')
        self.invoice = Invoice.objects.create(
            vendor=self.vendor,
            nomor='RACE-BS',
            tanggal=self.today,
            total_rp=10000,
            dibuat_oleh=self.buyer,
        )
        self.allocation = Alokasi.objects.create(
            cmt=self.cmt, status='disetujui', dibuat_oleh=self.buyer
        )
        self.roll = Roll.objects.create(
            invoice=self.invoice,
            material=self.material,
            color=color,
            yard=100,
            urut=1,
            status='diterima',
            alokasi=self.allocation,
            tgl_terima=self.today,
        )

    def race(self, operation):
        barrier = Barrier(2)

        def worker():
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                operation()
                return 'ok'
            except ValidationError:
                return 'rejected'
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as executor:
            return list(executor.map(lambda _: worker(), range(2)))

    def test_two_bs_requests_cannot_reserve_the_same_yards(self):
        results = self.race(
            lambda: record_defect(
                self.cmt,
                [{'roll': self.roll.pk, 'yard': '60'}],
                self.today,
                'Cacat',
                uuid4().hex,
                self.buyer,
            )
        )
        self.assertCountEqual(results, ['ok', 'rejected'])
        self.roll.refresh_from_db()
        self.assertEqual(self.roll.bs_yard, 60)
        self.assertEqual(FabricDefectItem.objects.count(), 1)

    def test_two_refund_receipts_cannot_exceed_invoice_refund(self):
        record_payment(self.invoice, payment_data(amount='10000', seen='10000'), self.director)
        record_defect(
            self.cmt,
            [{'roll': self.roll.pk, 'yard': '100'}],
            self.today,
            'Cacat',
            uuid4().hex,
            self.buyer,
        )
        results = self.race(
            lambda: record_refund(
                self.invoice,
                '7000',
                self.today,
                '',
                uuid4().hex,
                self.director,
            )
        )
        self.assertCountEqual(results, ['ok', 'rejected'])
        self.assertEqual(payment_summary(self.invoice).refund_received, 7000)
